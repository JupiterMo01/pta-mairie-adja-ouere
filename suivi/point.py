"""
suivi/point.py

Point des activités : liste des activités du PTA avec leur taux d'exécution physique,
leur statut et les observations cumulées de leurs tâches. Lecture seule : tout provient
du module Suivi & Évaluation, avec exactement les mêmes calculs (taux pondéré et statut
agrégé sur les tâches du périmètre consulté : global, direction ou service).
"""
import io
from datetime import datetime

from flask import render_template, request, redirect, url_for, flash, send_file
from flask_login import login_required, current_user

from models import db, Direction, Service
from suivi import suivi_bp
from suivi.routes import (_compute_pta_global, _compute_pta_direction, _compute_pta_service,
                          _filter_by_nature, _load_suivis_global, _enrich, _fmt_periode)
from utils import get_annee

STATUTS = {'execute': 'Exécutée', 'en_cours': 'En cours', 'non_execute': 'Non exécutée'}
STATUTS_PLURIEL = {'execute': 'Exécutées', 'en_cours': 'En cours', 'non_execute': 'Non exécutées'}
NATURES = {'fct': 'Fonctionnement', 'inv': 'Investissement'}


def _perimetre(annee):
    """Détermine le périmètre consulté selon le rôle (mêmes règles que le Suivi).
    Retourne un dict, ou None si le compte n'est rattaché à aucune entité."""
    role = current_user.role
    p = {'role': role, 'niveau': 'global', 'titre': '', 'data': [],
         'directions': [], 'services': [], 'services_dir': [],
         'sel_dir_id': None, 'sel_svc_id': None}

    if role == 'service':
        service = current_user.service
        if not service:
            return None
        p.update(data=_compute_pta_service(annee, service), niveau='service',
                 titre=f"{service.code} · {service.nom}")

    elif role == 'direction':
        direction = current_user.direction
        if not direction:
            return None
        p['services_dir'] = Service.query.filter_by(direction_id=direction.id).order_by(Service.nom).all()
        sel_svc_id = request.args.get('service_id', type=int)
        service = next((s for s in p['services_dir'] if s.id == sel_svc_id), None)
        if service:
            p.update(data=_compute_pta_service(annee, service), niveau='service',
                     titre=f"{service.code} · {service.nom}", sel_svc_id=service.id)
        else:
            p.update(data=_compute_pta_direction(annee, direction), niveau='direction',
                     titre=f"{direction.code} · {direction.nom}")

    else:  # admin_editeur, admin_lecteur
        p['directions'] = Direction.query.order_by(Direction.nom).all()
        p['services'] = Service.query.order_by(Service.nom).all()
        sel_svc_id = request.args.get('service_id', type=int)
        sel_dir_id = request.args.get('direction_id', type=int)
        service = db.session.get(Service, sel_svc_id) if sel_svc_id else None
        direction = db.session.get(Direction, sel_dir_id) if sel_dir_id and not service else None
        if service:
            p.update(data=_compute_pta_service(annee, service), niveau='service',
                     titre=f"{service.code} · {service.nom}", sel_svc_id=service.id)
        elif direction:
            p.update(data=_compute_pta_direction(annee, direction), niveau='direction',
                     titre=f"{direction.code} · {direction.nom}", sel_dir_id=direction.id)
        else:
            p.update(data=_compute_pta_global(annee), titre="Vue globale de tout le PTA")
    return p


def _filtres():
    nature = request.args.get('nature', '')
    statut = request.args.get('statut', '')
    return (nature if nature in NATURES else ''), (statut if statut in STATUTS else '')


def _observations(ad):
    """Observations/difficultés de toutes les tâches de l'activité, séparées par « ; »."""
    morceaux = []
    for td in ad['taches']:
        sv = td.get('suivi')
        texte = ' '.join((sv.observation or '').split()) if sv else ''
        if texte:
            morceaux.append(texte.rstrip(' ;.'))
    return ' ; '.join(morceaux)


def _construire(annee, perimetre, nature, statut):
    """Retourne (lignes affichées, compteurs par statut) pour le périmètre et les filtres."""
    data = _filter_by_nature(perimetre['data'], nature) if nature else perimetre['data']
    _enrich(data, _load_suivis_global(annee.id), None)

    toutes = []
    for pd in data:
        for pjd in pd['projets']:
            for ad in pjd['activites']:
                act = ad['activite']
                d = act.direction_responsable
                toutes.append({
                    'code_pta': ad['code'],
                    'libelle': act.nom,
                    'periode': _fmt_periode(act.periode_debut, act.periode_fin),
                    'direction_code': d.code if d else '',
                    'direction_nom': d.nom if d else '',
                    'taux': ad['taux'],
                    'statut': ad['statut'],
                    'observations': _observations(ad),
                })
    compteurs = {k: sum(1 for l in toutes if l['statut'] == k) for k in STATUTS}
    compteurs['tout'] = len(toutes)
    lignes = [l for l in toutes if not statut or l['statut'] == statut]
    for i, l in enumerate(lignes, 1):
        l['num'] = i
    return lignes, compteurs


def _libelle_filtres(nature, statut):
    morceaux = [NATURES[nature] if nature else 'Toutes natures',
                STATUTS_PLURIEL[statut] if statut else 'Tous statuts']
    return ' · '.join(morceaux)


def _preparer():
    annee = get_annee()
    if not annee:
        flash("Aucune année PTA active.", 'danger')
        return None, redirect(url_for('pta.global_pta'))
    perimetre = _perimetre(annee)
    if perimetre is None:
        flash("Votre compte n'est rattaché à aucune direction ni aucun service.", 'danger')
        return None, redirect(url_for('dashboard.index'))
    nature, statut = _filtres()
    lignes, compteurs = _construire(annee, perimetre, nature, statut)
    return dict(annee=annee, perimetre=perimetre, nature=nature, statut=statut,
                lignes=lignes, compteurs=compteurs,
                libelle_filtres=_libelle_filtres(nature, statut)), None


def _params_courants(**changes):
    """Paramètres d'URL actuels (périmètre et filtres), avec modifications."""
    q = {k: request.args.get(k) for k in ('direction_id', 'service_id', 'nature', 'statut')}
    q.update(changes)
    return {k: v for k, v in q.items() if v}


@suivi_bp.route('/point-activites')
@login_required
def point_activites():
    ctx, redir = _preparer()
    if redir:
        return redir
    return render_template('suivi/point.html', statuts=STATUTS, natures=NATURES,
                           params=_params_courants, **ctx)


@suivi_bp.route('/point-activites/imprimer')
@login_required
def point_imprimer():
    ctx, redir = _preparer()
    if redir:
        return redir
    return render_template('suivi/point_print.html', statuts=STATUTS, **ctx)


@suivi_bp.route('/point-activites/excel')
@login_required
def point_excel():
    ctx, redir = _preparer()
    if redir:
        return redir
    wb = _classeur(**ctx)
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    entite = ctx['perimetre']['titre'].split(' · ')[0].replace(' ', '_') if ctx['perimetre']['niveau'] != 'global' else 'Global'
    suffixe = '_'.join(filter(None, [NATURES.get(ctx['nature']), STATUTS_PLURIEL.get(ctx['statut'])]))
    nom = f"Point_des_activites_{ctx['annee'].annee}_{entite}{'_' + suffixe if suffixe else ''}.xlsx"
    nom = nom.replace('é', 'e').replace('É', 'E').replace('è', 'e')
    return send_file(tampon, as_attachment=True, download_name=nom,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _classeur(annee, perimetre, nature, statut, lignes, compteurs, libelle_filtres):
    """Classeur Excel du point des activités, avec en-tête institutionnel."""
    import os
    from flask import current_app
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Point des activités'
    NC = 7
    der = get_column_letter(NC)
    fin = Side(style='thin', color='7F8C85')
    moyen = Side(style='medium', color='155724')
    bord = Border(left=fin, right=fin, top=fin, bottom=fin)
    ctr = Alignment(horizontal='center', vertical='center', wrap_text=True)
    gch = Alignment(horizontal='left', vertical='center', wrap_text=True)
    dte = Alignment(horizontal='right', vertical='center', wrap_text=True)

    # Lignes 1 à 3 : bandeau, coordonnées, logo
    for r, h in ((1, 22), (2, 8), (3, 22)):
        ws.row_dimensions[r].height = h
    ws.merge_cells('A1:B3'); ws.merge_cells('C1:E3'); ws.merge_cells('F1:G3')
    c = ws['C1']
    c.value = "BP 02 Adja-Ouèrè\nTél : +229 01 61 91 96 12\nEmail : contact.adjaouere@mairie.bj"
    c.font = Font(bold=True, size=9)
    c.alignment = ctr
    try:
        from openpyxl.drawing.image import Image as XLImage
        from PIL import Image as PILImage
        dossier = os.path.join(current_app.root_path, 'static', 'img')
        hauteur = 52
        for fichier, ancre in (('bandeau.png', 'A1'), ('logo_commune.png', 'G1')):
            chemin = os.path.join(dossier, fichier)
            if not os.path.exists(chemin):
                continue
            tampon = io.BytesIO()
            with PILImage.open(chemin) as im:
                larg, haut = im.size
                im.save(tampon, 'PNG')
            tampon.seek(0)
            img = XLImage(tampon)
            img.height, img.width = hauteur, int(larg * hauteur / haut)
            ws.add_image(img, ancre)
    except Exception:
        pass

    # Titre et sous-titre
    titre_fill = PatternFill('solid', fgColor='D1F0DA')
    ws.merge_cells(f'A4:{der}4')
    ws['A4'].value = f"POINT DES ACTIVITÉS DU PTA · Exercice {annee.annee}"
    ws['A4'].font = Font(bold=True, size=13, color='0F3529')
    ws['A4'].alignment = ctr
    ws.merge_cells(f'A5:{der}5')
    ws['A5'].value = f"{perimetre['titre']}   |   {libelle_filtres}"
    ws['A5'].font = Font(bold=True, size=10, color='155724')
    ws['A5'].alignment = ctr
    ws.merge_cells(f'A6:{der}6')
    ws['A6'].value = (f"Activités : {compteurs['tout']}   |   Exécutées : {compteurs['execute']}   |   "
                      f"En cours : {compteurs['en_cours']}   |   Non exécutées : {compteurs['non_execute']}")
    ws['A6'].font = Font(size=9, color='333333')
    ws['A6'].alignment = ctr
    for r in (4, 5, 6):
        ws.row_dimensions[r].height = 24 if r == 4 else 18
        for col in range(1, NC + 1):
            cell = ws.cell(row=r, column=col)
            cell.fill = titre_fill
            cell.border = Border(left=moyen if col == 1 else None, right=moyen if col == NC else None,
                                 top=moyen if r == 4 else None, bottom=moyen if r == 6 else None)

    # En-têtes
    ent = 8
    entetes = ['N°', 'Activité', "Période d'exécution", 'Direction responsable',
               "Taux d'exécution physique", 'Statut', 'Observations / Difficultés']
    for col, h in enumerate(entetes, 1):
        c = ws.cell(row=ent, column=col, value=h)
        c.font = Font(bold=True, size=9, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='0F3529')
        c.alignment = ctr
        c.border = bord
    ws.row_dimensions[ent].height = 30

    couleurs = {'execute': ('D4EDDA', '155724'), 'en_cours': ('FFF3CD', '856404'),
                'non_execute': ('E9ECEF', '495057')}
    r = ent
    for l in lignes:
        r += 1
        fond = PatternFill('solid', fgColor='F7FAF8') if l['num'] % 2 == 0 else None
        valeurs = [l['num'], l['libelle'], l['periode'], l['direction_code'],
                   l['taux'] / 100, STATUTS[l['statut']], l['observations']]
        for col, v in enumerate(valeurs, 1):
            c = ws.cell(row=r, column=col, value=v)
            c.font = Font(size=9)
            c.border = bord
            c.alignment = gch if col in (2, 7) else ctr
            if fond:
                c.fill = fond
        ws.cell(row=r, column=5).number_format = '0.00%'
        f_st, t_st = couleurs[l['statut']]
        cs = ws.cell(row=r, column=6)
        cs.fill = PatternFill('solid', fgColor=f_st)
        cs.font = Font(size=9, bold=True, color=t_st)
    if not lignes:
        r += 1
        ws.merge_cells(f'A{r}:{der}{r}')
        ws.cell(row=r, column=1, value="Aucune activité pour cette sélection.").alignment = ctr

    # Pied
    r += 2
    ws.merge_cells(f'A{r}:C{r}')
    c = ws.cell(row=r, column=1, value=f"Exporté le {datetime.now().strftime('%d/%m/%Y à %H:%M')}")
    c.font = Font(italic=True, size=8, color='555555')
    c.alignment = gch
    ws.merge_cells(f'D{r}:{der}{r}')
    c = ws.cell(row=r, column=4, value="Direction du Développement Local et de la Planification (DDLP)")
    c.font = Font(bold=True, size=8)
    c.alignment = dte

    for col, larg in enumerate([6, 52, 16, 14, 14, 15, 60], 1):
        ws.column_dimensions[get_column_letter(col)].width = larg
    ws.freeze_panes = f'A{ent + 1}'
    if lignes:
        ws.auto_filter.ref = f'A{ent}:{der}{ent + len(lignes)}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{ent}:{ent}'
    ws.oddFooter.left.text = f"Point des activités du PTA {annee.annee}"
    ws.oddFooter.right.text = "Page &P / &N"
    return wb
