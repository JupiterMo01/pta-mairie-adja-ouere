import io
from flask import render_template, abort, request, jsonify, Response
from flask_login import login_required, current_user
from models import db, Programme, Projet, Activite, PaiActivite, PaiProgramme, PaiProjet
from pai import pai_bp
from utils import get_annee


# ─── Helper commun ────────────────────────────────────────────────────────────

def _build_pai_data(annee):
    """Construit la structure PAI depuis les activités d'investissement marquées inclure_dans_pai."""
    programmes = (
        Programme.query
        .filter_by(annee_id=annee.id)
        .order_by(Programme.numero)
        .all()
    )
    pai_data = []
    total_fp = total_fadec = total_ptfs = total_global = 0.0
    total_nb = 0
    prog_num = 0

    for pg in programmes:
        pg_projets = []
        proj_num = 0
        for pj in pg.projets:
            inv_acts = [a for a in pj.activites
                        if a.type_activite and 'investissement' in a.type_activite.lower()
                        and a.inclure_dans_pai is not False]
            if inv_acts:
                proj_num += 1
                for a in inv_acts:
                    total_fp     += a.src_rp
                    total_fadec  += a.src_fa + a.src_fn
                    total_ptfs   += a.src_ap + a.src_af
                    total_global += a.budget_total
                    total_nb     += 1
                pg_projets.append({'projet': pj, 'activites': inv_acts, 'proj_num': proj_num})
        if pg_projets:
            prog_num += 1
            pai_data.append({'programme': pg, 'projets': pg_projets, 'prog_num': prog_num})

    return pai_data, total_fp, total_fadec, total_ptfs, total_global, total_nb


# ─── Routes ──────────────────────────────────────────────────────────────────

@pai_bp.route('/')
@login_required
def index():
    """Vue écran du PAI (étend base.html). Admins seulement."""
    if current_user.role not in ('admin_editeur', 'admin_lecteur'):
        abort(403)
    annee = get_annee()
    if not annee:
        abort(404)

    pai_data, total_fp, total_fadec, total_ptfs, total_global, total_nb = _build_pai_data(annee)

    return render_template(
        'pai/index.html',
        annee=annee,
        pai_data=pai_data,
        total_fp=total_fp,
        total_fadec=total_fadec,
        total_ptfs=total_ptfs,
        total_global=total_global,
        total_nb=total_nb,
        can_edit=(current_user.role == 'admin_editeur'),
    )


@pai_bp.route('/print')
@login_required
def print_view():
    """Vue impression A4 paysage du PAI. Admins seulement."""
    if current_user.role not in ('admin_editeur', 'admin_lecteur'):
        abort(403)
    annee = get_annee()
    if not annee:
        abort(404)

    pai_data, total_fp, total_fadec, total_ptfs, total_global, total_nb = _build_pai_data(annee)

    return render_template(
        'pai/print.html',
        annee=annee,
        pai_data=pai_data,
        total_fp=total_fp,
        total_fadec=total_fadec,
        total_ptfs=total_ptfs,
        total_global=total_global,
    )


@pai_bp.route('/export')
@login_required
def export_excel():
    """Exporte le PAI en fichier Excel (.xlsx). Admins seulement."""
    if current_user.role not in ('admin_editeur', 'admin_lecteur'):
        abort(403)
    annee = get_annee()
    if not annee:
        abort(404)

    pai_data, total_fp, total_fadec, total_ptfs, total_global, _ = _build_pai_data(annee)

    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f'PAI {annee.annee}'

    def fill(hex_color):
        return PatternFill('solid', fgColor=hex_color)

    thin = Side(style='thin')
    bord = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    right  = Alignment(horizontal='right', vertical='center')
    left   = Alignment(horizontal='left',  vertical='center', wrap_text=True)

    headers = [
        'Code PAI', 'Activités', 'Localisation', 'Poids (%)', 'Indicateurs',
        "Période d'exécution", 'Struct. Resp.', 'Structures associées',
        'FP (milliers F CFA)', 'FADeC (libellé)', 'Montant FADeC (milliers)',
        'Autres PTFs (milliers)', 'Coût Total (milliers)', 'Observations',
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True, size=9)
        c.fill = fill('BDD7EE')
        c.alignment = center
        c.border = bord
    ws.row_dimensions[1].height = 30

    def _periode(act):
        if act.periode_debut and act.periode_fin and act.periode_debut != act.periode_fin:
            return f'{act.periode_debut} – {act.periode_fin}'
        return act.periode_debut or act.periode_fin or ''

    def _assoc(act):
        codes = [d.code for d in act.directions_associees] + [s.code for s in act.services_intervenants]
        return ', '.join(codes)

    r = 2
    for pg_d in pai_data:
        pg = pg_d['programme']
        pn = pg_d['prog_num']
        extra_pg = pg.pai_extra_prog
        poids_pg = extra_pg.poids_pai if extra_pg else ''
        ws.cell(row=r, column=1, value=str(pn))
        ws.cell(row=r, column=2, value=f'Programme {pn} : {pg.nom}')
        ws.cell(row=r, column=4, value=poids_pg if poids_pg else '')
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=14)
        for col in range(1, 15):
            c = ws.cell(row=r, column=col)
            c.fill = fill('F4B183'); c.font = Font(bold=True, size=9); c.border = bord
        ws.cell(row=r, column=1).alignment = center
        ws.cell(row=r, column=2).alignment = left
        ws.cell(row=r, column=4).alignment = center
        r += 1

        for pj_d in pg_d['projets']:
            pj = pj_d['projet']
            pjn = pj_d['proj_num']
            extra_pj = pj.pai_extra_proj
            poids_pj = extra_pj.poids_pai if extra_pj else ''
            ws.cell(row=r, column=1, value=f'{pn}.{pjn}')
            ws.cell(row=r, column=2, value=f'Projet {pn}.{pjn} : {pj.nom}')
            ws.cell(row=r, column=4, value=poids_pj if poids_pj else '')
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=14)
            for col in range(1, 15):
                c = ws.cell(row=r, column=col)
                c.fill = fill('FFFF00'); c.font = Font(bold=True, size=9); c.border = bord
            ws.cell(row=r, column=1).alignment = center
            ws.cell(row=r, column=2).alignment = left
            ws.cell(row=r, column=4).alignment = center
            r += 1

            pj_fp = pj_fadec = pj_ptfs = pj_total = 0.0

            for act_i, act in enumerate(pj_d['activites'], 1):
                extra = act.pai_extra
                a_fp    = act.src_rp / 1000
                a_fadec = (act.src_fa + act.src_fn) / 1000
                a_ptfs  = (act.src_ap + act.src_af) / 1000
                a_total = act.budget_total / 1000
                pj_fp += a_fp; pj_fadec += a_fadec; pj_ptfs += a_ptfs; pj_total += a_total

                row_data = [
                    f'{pn}.{pjn}.{act_i}',
                    act.nom,
                    extra.localisation     if extra and extra.localisation     else '',
                    extra.poids_pai        if extra and extra.poids_pai        else '',
                    extra.indicateurs      if extra and extra.indicateurs      else '',
                    _periode(act),
                    act.direction_responsable.code if act.direction_responsable else '',
                    _assoc(act),
                    a_fp    or '',
                    extra.fadec_type       if extra and extra.fadec_type       else '',
                    a_fadec or '',
                    a_ptfs  or '',
                    a_total or '',
                    extra.observations_pai if extra and extra.observations_pai else '',
                ]
                for col, val in enumerate(row_data, 1):
                    c = ws.cell(row=r, column=col, value=val)
                    c.fill = fill('C5DEB5'); c.font = Font(size=9); c.border = bord
                    if col in (9, 11, 12, 13):
                        c.number_format = '#,##0.000'; c.alignment = right
                    elif col in (1, 4, 6, 7):
                        c.alignment = center
                    else:
                        c.alignment = left
                r += 1

            st = ['', f'Sous-Total {pn}.{pjn}', '', '', '', '', '', '',
                  pj_fp, '', pj_fadec, pj_ptfs, pj_total, '']
            for col, val in enumerate(st, 1):
                c = ws.cell(row=r, column=col, value=val)
                c.fill = fill('EAF0FB'); c.font = Font(bold=True, italic=True, size=9); c.border = bord
                if col in (9, 11, 12, 13):
                    c.number_format = '#,##0.000'; c.alignment = right
                elif col == 2:
                    c.alignment = Alignment(horizontal='right', vertical='center')
            r += 1

    tot = ['', f'TOTAL PAI {annee.annee}', '', '', '', '', '', '',
           total_fp / 1000, '', total_fadec / 1000, total_ptfs / 1000, total_global / 1000, '']
    for col, val in enumerate(tot, 1):
        c = ws.cell(row=r, column=col, value=val)
        c.fill = fill('1A3A5C'); c.font = Font(bold=True, size=9, color='FFFFFF'); c.border = bord
        if col in (9, 11, 12, 13):
            c.number_format = '#,##0.000'; c.alignment = right
        elif col == 2:
            c.alignment = Alignment(horizontal='right', vertical='center')

    for i, w in enumerate([8, 35, 18, 8, 25, 12, 12, 20, 16, 22, 16, 16, 16, 22], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        output.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename=PAI_{annee.annee}.xlsx'},
    )


@pai_bp.route('/edit/<int:activite_id>', methods=['POST'])
@login_required
def edit(activite_id):
    """Sauvegarde les champs PAI-spécifiques d'une activité (admin éditeur uniquement)."""
    if current_user.role != 'admin_editeur':
        abort(403)

    activite = db.session.get(Activite, activite_id)
    if not activite:
        abort(404)

    extra = PaiActivite.query.filter_by(activite_id=activite_id).first()
    if not extra:
        extra = PaiActivite(activite_id=activite_id)
        # Importer le poids depuis le PTA si non encore défini
        extra.poids_pai = activite.poids or 0.0
        db.session.add(extra)

    extra.localisation     = request.form.get('localisation', '').strip() or None
    poids_val = request.form.get('poids_pai', '').strip()
    extra.poids_pai        = float(poids_val) if poids_val else (activite.poids or 0.0)
    extra.indicateurs      = request.form.get('indicateurs', '').strip() or None
    extra.fadec_type       = request.form.get('fadec_type', '').strip() or None
    extra.observations_pai = request.form.get('observations_pai', '').strip() or None

    db.session.commit()
    return jsonify({'ok': True})


@pai_bp.route('/edit_prog/<int:programme_id>', methods=['POST'])
@login_required
def edit_prog(programme_id):
    """Sauvegarde le poids PAI d'un programme."""
    if current_user.role != 'admin_editeur':
        abort(403)

    programme = db.session.get(Programme, programme_id)
    if not programme:
        abort(404)

    extra = PaiProgramme.query.filter_by(programme_id=programme_id).first()
    if not extra:
        extra = PaiProgramme(programme_id=programme_id)
        db.session.add(extra)

    poids_val = request.form.get('poids_pai', '').strip()
    extra.poids_pai = float(poids_val) if poids_val else 0.0

    db.session.commit()
    return jsonify({'ok': True})


@pai_bp.route('/edit_proj/<int:projet_id>', methods=['POST'])
@login_required
def edit_proj(projet_id):
    """Sauvegarde le poids PAI d'un projet."""
    if current_user.role != 'admin_editeur':
        abort(403)

    projet = db.session.get(Projet, projet_id)
    if not projet:
        abort(404)

    extra = PaiProjet.query.filter_by(projet_id=projet_id).first()
    if not extra:
        extra = PaiProjet(projet_id=projet_id)
        db.session.add(extra)

    poids_val = request.form.get('poids_pai', '').strip()
    extra.poids_pai = float(poids_val) if poids_val else 0.0

    db.session.commit()
    return jsonify({'ok': True})
