import io
import os
import datetime
from functools import wraps
from flask import render_template, redirect, url_for, flash, send_file, current_app
from flask_login import login_required, current_user
from statspai import statspai_bp
from utils import get_annee, MOIS_ORDRE as MOIS_NUM
from pai.routes import _build_pai_data


def admin_only(f):
    @wraps(f)
    @login_required
    def decorated(*args, **kwargs):
        if current_user.role not in ('admin_editeur', 'admin_lecteur'):
            flash('Accès réservé aux administrateurs.', 'danger')
            return redirect(url_for('pta.global_pta'))
        return f(*args, **kwargs)
    return decorated


END_TRIM = {1: 3, 2: 6, 3: 9, 4: 12}


def _fmt(v):
    try:
        return '{:,.0f}'.format(float(v or 0)).replace(',', ' ')
    except (TypeError, ValueError):
        return '0'


def _fr(v, d=2):
    try:
        return ('{:.' + str(d) + 'f}').format(float(v or 0)).replace('.', ',')
    except (TypeError, ValueError):
        return '0,' + '0' * d


# ── Contribution trimestrielle d'une activité ─────────────────────────────────

def _act_contrib(act, trim):
    """Proportion (0.0–1.0) d'une activité réalisée à fin du trimestre `trim`.
    Pro-ratisation linéaire mensuelle sur [periode_debut, periode_fin].
    Si la période n'est pas renseignée : 25 %/50 %/75 %/100 %.
    """
    m_d = MOIS_NUM.get(act.periode_debut or '', 1)
    m_f = MOIS_NUM.get(act.periode_fin   or '', 12)
    if m_f < m_d:
        m_f = m_d
    duree     = m_f - m_d + 1
    completed = min(m_f, END_TRIM[trim]) - m_d + 1
    return max(0.0, min(completed / duree, 1.0))


# ── Calcul des statistiques PAI ───────────────────────────────────────────────

def _src5(acts):
    """Retourne (rp, fa, fn, ap, af, total) pour une liste d'activités."""
    rp = sum(a.src_rp for a in acts)
    fa = sum(a.src_fa for a in acts)
    fn = sum(a.src_fn for a in acts)
    ap = sum(a.src_ap for a in acts)
    af = sum(a.src_af for a in acts)
    return rp, fa, fn, ap, af, rp + fa + fn + ap + af


def _compute_stats_pai(annee):
    pai_data, _fp, _fadec, _ptfs, total_global, total_nb = _build_pai_data(annee)

    all_acts = [a for pg_d in pai_data
                  for pj_d in pg_d['projets']
                  for a in pj_d['activites']]
    total_rp, total_fa, total_fn, total_ap, total_af, _ = _src5(all_acts)

    def pct(val):
        return round(val / total_global * 100, 2) if total_global else 0

    # Stats par programme et projet
    stats_programmes = []
    for pg_d in pai_data:
        pg      = pg_d['programme']
        acts_pg = [a for pj_d in pg_d['projets'] for a in pj_d['activites']]
        rp_pg, fa_pg, fn_pg, ap_pg, af_pg, bud_pg = _src5(acts_pg)
        projets_data = []
        for pj_d in pg_d['projets']:
            pj = pj_d['projet']
            acts_pj = pj_d['activites']
            rp_pj, fa_pj, fn_pj, ap_pj, af_pj, bud_pj = _src5(acts_pj)
            projets_data.append({
                'code':         f"{pg_d['prog_num']}.{pj_d['proj_num']}",
                'nom':          pj.nom,
                'poids':        pj_d['poids'],
                'nb_activites': len(acts_pj),
                'rp': rp_pj, 'fa': fa_pj, 'fn': fn_pj,
                'ap': ap_pj, 'af': af_pj, 'budget': bud_pj,
                'fadec': fa_pj + fn_pj, 'ptfs': ap_pj + af_pj,
            })
        stats_programmes.append({
            'num':  pg_d['prog_num'],
            'code': pg.code,
            'nom':  pg.nom,
            'poids': pg_d['poids'],
            'nb_projets':   len(pg_d['projets']),
            'nb_activites': len(acts_pg),
            'rp': rp_pg, 'fa': fa_pg, 'fn': fn_pg,
            'ap': ap_pg, 'af': af_pg, 'budget': bud_pg,
            'fadec': fa_pg + fn_pg, 'ptfs': ap_pg + af_pg,
            'projets': projets_data,
        })

    # Stats par direction
    dir_map = {}
    for pg_d in pai_data:
        for pj_d in pg_d['projets']:
            for act in pj_d['activites']:
                d   = act.direction_responsable
                key = d.id if d else 0
                if key not in dir_map:
                    dir_map[key] = {
                        'code': d.code if d else '—',
                        'nom':  d.nom  if d else '—',
                        'nb_activites': 0,
                        'rp': 0.0, 'fa': 0.0, 'fn': 0.0,
                        'ap': 0.0, 'af': 0.0, 'budget': 0.0,
                    }
                dir_map[key]['nb_activites'] += 1
                dir_map[key]['rp'] += act.src_rp
                dir_map[key]['fa'] += act.src_fa
                dir_map[key]['fn'] += act.src_fn
                dir_map[key]['ap'] += act.src_ap
                dir_map[key]['af'] += act.src_af
                dir_map[key]['budget'] += act.budget_total
    for d in dir_map.values():
        d['fadec'] = d['fa'] + d['fn']
        d['ptfs']  = d['ap'] + d['af']
        d['fp']    = d['rp']             # alias rétrocompatible
    stats_directions = sorted(dir_map.values(), key=lambda x: x['code'])

    # Liste détaillée des activités
    liste_activites = []
    for pg_d in pai_data:
        pn = pg_d['prog_num']
        for pj_d in pg_d['projets']:
            pjn = pj_d['proj_num']
            for act_i, act in enumerate(pj_d['activites'], 1):
                extra = act.pai_extra
                liste_activites.append({
                    'code_pai':    f'{pn}.{pjn}.{act_i}',
                    'nom':         act.nom,
                    'direction':   act.direction_responsable.code if act.direction_responsable else '—',
                    'periode':     (f"{act.periode_debut}–{act.periode_fin}"
                                    if act.periode_debut and act.periode_fin
                                    else act.periode_debut or act.periode_fin or '—'),
                    'poids':       int(extra.poids_pai or 0) if extra else 0,
                    'rp': act.src_rp, 'fa': act.src_fa, 'fn': act.src_fn,
                    'ap': act.src_ap, 'af': act.src_af,
                    'budget':      act.budget_total,
                    'localisation': extra.localisation if extra else '',
                    'indicateurs':  extra.indicateurs  if extra else '',
                })

    return dict(
        pai_data=pai_data,
        total_rp=total_rp, total_fa=total_fa, total_fn=total_fn,
        total_ap=total_ap, total_af=total_af,
        total_fp=total_rp,                            # alias RP = FP
        total_fadec=total_fa + total_fn,              # regroupé pour tableaux
        total_ptfs=total_ap + total_af,               # regroupé pour tableaux
        total_budget=total_global,
        nb_programmes=len(pai_data),
        nb_projets=sum(len(pg_d['projets']) for pg_d in pai_data),
        nb_activites=total_nb,
        stats_programmes=stats_programmes,
        stats_directions=stats_directions,
        liste_activites=liste_activites,
        pct=pct,
    )


# ── Cibles trimestrielles PAI ─────────────────────────────────────────────────

def _compute_cibles_pai(pai_data):
    """Cibles par programme/projet calculées depuis les activités (poids_pai + période)."""
    all_prog = []

    for pg_d in pai_data:
        proj_entries = []
        for pj_d in pg_d['projets']:
            act_data = []
            for act in pj_d['activites']:
                extra = act.pai_extra
                poids = float(extra.poids_pai or 0) if extra else 0.0
                cibles_act = {trim: _act_contrib(act, trim) * 100 for trim in (1, 2, 3, 4)}
                act_data.append((poids, cibles_act))

            if not act_data:
                continue
            tot_a = sum(p for p, _ in act_data) or len(act_data)
            proj_cibles = {
                trim: round(sum(p * c[trim] for p, c in act_data) / tot_a, 1)
                for trim in (1, 2, 3, 4)
            }
            proj_entries.append({
                'code':   f"{pg_d['prog_num']}.{pj_d['proj_num']}",
                'nom':    pj_d['projet'].nom,
                'poids':  pj_d['poids'],
                'cibles': proj_cibles,
            })

        if not proj_entries:
            continue
        tot_pr = sum(p['poids'] for p in proj_entries) or len(proj_entries)
        prog_cibles = {
            trim: round(sum(p['cibles'][trim] * p['poids'] for p in proj_entries) / tot_pr, 1)
            for trim in (1, 2, 3, 4)
        }
        all_prog.append({
            'code':    str(pg_d['prog_num']),
            'nom':     pg_d['programme'].nom,
            'poids':   pg_d['poids'],
            'cibles':  prog_cibles,
            'projets': proj_entries,
        })

    if not all_prog:
        return {'global': {1: 0, 2: 0, 3: 0, 4: 0}, 'programmes': []}
    tot_pg = sum(p['poids'] for p in all_prog) or len(all_prog)
    global_cibles = {
        trim: round(sum(p['cibles'][trim] * p['poids'] for p in all_prog) / tot_pg, 1)
        for trim in (1, 2, 3, 4)
    }
    return {'global': global_cibles, 'programmes': all_prog}


def _compute_cibles_directions_pai(pai_data):
    """Cibles trimestrielles par direction responsable (PAI)."""
    dir_acts = {}
    for pg_d in pai_data:
        for pj_d in pg_d['projets']:
            for act in pj_d['activites']:
                extra = act.pai_extra
                poids = float(extra.poids_pai or 0) if extra else 0.0
                d    = act.direction_responsable
                key  = d.code if d else '—'
                nom  = d.nom  if d else '—'
                cibles_act = {trim: _act_contrib(act, trim) * 100 for trim in (1, 2, 3, 4)}
                if key not in dir_acts:
                    dir_acts[key] = {'nom': nom, 'acts': []}
                dir_acts[key]['acts'].append((poids, cibles_act))

    result = []
    for code, data in sorted(dir_acts.items()):
        acts  = data['acts']
        tot_p = sum(p for p, _ in acts) or len(acts)
        cibles = {
            trim: round(sum(p * c[trim] for p, c in acts) / tot_p, 1)
            for trim in (1, 2, 3, 4)
        }
        result.append({
            'code': code, 'nom': data['nom'],
            'nb_activites': len(acts), 'cibles': cibles,
        })
    return result


# ── Graphiques ────────────────────────────────────────────────────────────────

def _chart_sources_pai(rp, fa, fn, ap, af, title='Répartition par source de financement'):
    try:
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt

        labels = ['Fonds Propres (RP)', 'FADeC Affecté (FA)', 'FADeC Non Affecté (FNA)',
                  'Appui Partenaires (AP)', 'Autres Financements (AF)']
        values = [rp, fa, fn, ap, af]
        colors = ['#1F4E79', '#2E86AB', '#3BB273', '#E9C46A', '#E76F51']

        pairs = [(l, v, c) for l, v, c in zip(labels, values, colors) if v > 0]
        if not pairs:
            return None
        f_labels, f_values, f_colors = zip(*pairs)

        fig, ax = plt.subplots(figsize=(7, 4.5), dpi=110)
        wedges, _, autotexts = ax.pie(
            f_values, labels=None, colors=f_colors,
            autopct='%1.1f%%', startangle=90,
            wedgeprops=dict(width=0.55, edgecolor='white', linewidth=1.5)
        )
        for at in autotexts:
            at.set_fontsize(9)
        ax.legend(wedges, f_labels, loc='lower center',
                  bbox_to_anchor=(0.5, -0.28), ncol=2, fontsize=8.5, frameon=False)
        ax.set_title(title, fontsize=11, fontweight='bold', pad=14)
        plt.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format='png', bbox_inches='tight', dpi=130)
        plt.close(fig)
        buf.seek(0)
        return buf
    except Exception:
        return None


def _chart_directions_pai(dirs_data):
    try:
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        import numpy as np

        actifs = [d for d in dirs_data if d['nb_activites'] > 0]
        if not actifs:
            return None
        labels = [d['code'] for d in actifs]
        n_act  = [d['nb_activites'] for d in actifs]

        fig_h = max(3.0, len(labels) * 0.5 + 1.2)
        fig, ax = plt.subplots(figsize=(9, fig_h), dpi=110)
        y = np.arange(len(labels))
        ax.barh(y, n_act, color='#1F4E79', alpha=0.85)
        ax.set_yticks(y)
        ax.set_yticklabels(labels, fontsize=9)
        ax.set_xlabel('Nombre d\'activités', fontsize=9)
        ax.set_title('Activités PAI par direction responsable', fontsize=11, fontweight='bold')
        ax.xaxis.set_major_locator(plt.MaxNLocator(integer=True))
        ax.invert_yaxis()
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        plt.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format='png', bbox_inches='tight', dpi=130)
        plt.close(fig)
        buf.seek(0)
        return buf
    except Exception:
        return None


# ── Rapport Word ──────────────────────────────────────────────────────────────

def _build_word_pai(annee, s, static_img_path, cibles=None, cibles_directions=None):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    cibles            = cibles or {}
    cibles_directions = cibles_directions or []

    def shade(cell, hex_col):
        tc   = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd  = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_col)
        tcPr.append(shd)

    def borders(table):
        tbl   = table._tbl
        tblPr = tbl.tblPr
        if tblPr is None:
            tblPr = OxmlElement('w:tblPr')
            tbl.insert(0, tblPr)
        tblBrd = OxmlElement('w:tblBorders')
        for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            b = OxmlElement(f'w:{side}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), '4')
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), 'AAAAAA')
            tblBrd.append(b)
        tblPr.append(tblBrd)

    def repeat_hdr(row):
        tr   = row._tr
        trPr = tr.get_or_add_trPr()
        hdr  = OxmlElement('w:tblHeader')
        hdr.set(qn('w:val'), 'true')
        trPr.append(hdr)

    def switch_landscape(doc):
        sec = doc.add_section()
        sec.orientation   = WD_ORIENT.LANDSCAPE
        sec.page_width    = Cm(29.7)
        sec.page_height   = Cm(21)
        sec.left_margin   = sec.right_margin  = Cm(1.5)
        sec.top_margin    = sec.bottom_margin = Cm(1.5)

    def switch_portrait(doc):
        sec = doc.add_section()
        sec.orientation   = WD_ORIENT.PORTRAIT
        sec.page_width    = Cm(21)
        sec.page_height   = Cm(29.7)
        sec.left_margin   = sec.right_margin  = Cm(2.5)
        sec.top_margin    = Cm(2)
        sec.bottom_margin = Cm(2)

    ALGN = {
        'left':   WD_ALIGN_PARAGRAPH.LEFT,
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'right':  WD_ALIGN_PARAGRAPH.RIGHT,
    }

    def h1(doc, text):
        p   = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after  = Pt(4)
        r   = p.add_run(text)
        r.bold = True; r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    def h2(doc, text, color=(0x2E, 0x86, 0xAB)):
        p   = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after  = Pt(2)
        r   = p.add_run(text)
        r.bold = True; r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(*color)

    def interp(doc, text):
        p   = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(8)
        r   = p.add_run('ℹ️  ' + text)
        r.italic = True; r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)

    def hdr_cell(cell, text, bg='1F4E79'):
        p   = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r   = p.add_run(text)
        r.bold = True; r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, bg)

    def data_cell(cell, text, align='left', bold=False, size=9, fg=None):
        p   = cell.paragraphs[0]
        p.alignment = ALGN.get(align, WD_ALIGN_PARAGRAPH.LEFT)
        r   = p.add_run(str(text) if text is not None else '')
        r.bold = bold; r.font.size = Pt(size)
        if fg:
            r.font.color.rgb = RGBColor(*fg)

    def cible_cell(cell, val):
        v = float(val) if val is not None else 0.0
        if   v < 25: col = 'FADBD8'
        elif v < 50: col = 'FDEBD0'
        elif v < 75: col = 'D6EAF8'
        else:        col = 'D5F5E3'
        data_cell(cell, f"{_fr(v)} %", align='center', bold=True, size=9)
        shade(cell, col)

    # ── Document ──────────────────────────────────────────────────────────────
    doc = Document()
    sec0 = doc.sections[0]
    sec0.orientation  = WD_ORIENT.PORTRAIT
    sec0.page_width   = Cm(21); sec0.page_height = Cm(29.7)
    sec0.left_margin  = sec0.right_margin = Cm(2.5)
    sec0.top_margin   = Cm(2.5); sec0.bottom_margin = Cm(2)

    # Page de titre
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('RAPPORT STATISTIQUE')
    r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f"PLAN ANNUEL D'INVESTISSEMENT (PAI) — Exercice {annee.annee}")
    r2.bold = True; r2.font.size = Pt(14); r2.font.color.rgb = RGBColor(0x2E, 0x86, 0xAB)

    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("Mairie d'Adja-Ouèrè")
    r3.bold = True; r3.font.size = Pt(12)

    logo_path = os.path.join(static_img_path, 'logo_commune.png')
    if os.path.exists(logo_path):
        pl = doc.add_paragraph(); pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pl.add_run().add_picture(logo_path, height=Cm(2.4))

    p4 = doc.add_paragraph(); p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = p4.add_run(f"Généré le {datetime.datetime.now().strftime('%d/%m/%Y à %H:%M')}")
    r4.italic = True; r4.font.size = Pt(9); r4.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # ════════════════════════════════════════════════════════
    # I. SYNTHÈSE GLOBALE
    # ════════════════════════════════════════════════════════
    h1(doc, 'I. SYNTHÈSE GLOBALE')

    tbl_kpi = doc.add_table(rows=2, cols=4)
    borders(tbl_kpi)
    repeat_hdr(tbl_kpi.rows[0])
    for i, (h_txt, v_txt) in enumerate(zip(
        ['Budget total PAI', 'Programmes', 'Projets', 'Activités d\'investissement'],
        [f"{_fmt(s['total_budget'])} F CFA",
         str(s['nb_programmes']), str(s['nb_projets']), str(s['nb_activites'])],
    )):
        hdr_cell(tbl_kpi.cell(0, i), h_txt)
        c = tbl_kpi.cell(1, i)
        data_cell(c, v_txt, align='center', bold=True, size=11)
        shade(c, 'D6EAF8')
    doc.add_paragraph()

    def _pct_s(v):
        return _fr(round(v / s['total_budget'] * 100, 2)) if s['total_budget'] else '0,00'
    interp(doc,
        f"Le PAI {annee.annee} comprend {s['nb_programmes']} programme(s) répartis en "
        f"{s['nb_projets']} projet(s) et {s['nb_activites']} activité(s) d'investissement. "
        f"Le budget global s'élève à {_fmt(s['total_budget'])} F CFA, dont {_pct_s(s['total_rp'])} % "
        f"sur Ressources Propres, {_pct_s(s['total_fa'])} % FADeC Affecté, "
        f"{_pct_s(s['total_fn'])} % FADeC Non Affecté, {_pct_s(s['total_ap'])} % "
        f"Appui Partenaires et {_pct_s(s['total_af'])} % Autres Financements."
    )

    # ════════════════════════════════════════════════════════
    # II. RÉPARTITION PAR SOURCE
    # ════════════════════════════════════════════════════════
    doc.add_page_break()
    h1(doc, 'II. RÉPARTITION BUDGÉTAIRE PAR SOURCE DE FINANCEMENT')

    SRC_DATA = [
        ('Ressources Propres (RP)',       s['total_rp'], 'EBF5FB'),
        ('FADeC Affecté (FA)',            s['total_fa'], 'D6EAF8'),
        ('FADeC Non Affecté (FNA)',       s['total_fn'], 'D5F5E3'),
        ('Appui Partenaires (AP)',        s['total_ap'], 'FEFDE7'),
        ('Autres Financements (AF)',      s['total_af'], 'FDEBD0'),
    ]
    tbl_src = doc.add_table(rows=len(SRC_DATA) + 2, cols=3)
    borders(tbl_src)
    repeat_hdr(tbl_src.rows[0])
    for i, txt in enumerate(['Source de financement', 'Montant (F CFA)', 'Part (%)']):
        hdr_cell(tbl_src.cell(0, i), txt, bg='2C3E50')
    for ri, (lbl, val, col) in enumerate(SRC_DATA, 1):
        data_cell(tbl_src.rows[ri].cells[0], lbl, size=9)
        data_cell(tbl_src.rows[ri].cells[1], _fmt(val), align='right', size=9)
        data_cell(tbl_src.rows[ri].cells[2], f"{_fr(s['pct'](val))} %", align='center', size=9)
        for cell in tbl_src.rows[ri].cells:
            shade(cell, col)
    tot_row = tbl_src.rows[-1]
    data_cell(tot_row.cells[0], 'TOTAL', bold=True, size=9)
    data_cell(tot_row.cells[1], _fmt(s['total_budget']), align='right', bold=True, size=9)
    data_cell(tot_row.cells[2], '100 %', align='center', bold=True, size=9)
    for cell in tot_row.cells:
        shade(cell, 'AED6F1')

    chart_src = _chart_sources_pai(s['total_rp'], s['total_fa'], s['total_fn'], s['total_ap'], s['total_af'])
    if chart_src:
        doc.add_picture(chart_src, width=Cm(12))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # ════════════════════════════════════════════════════════
    # III. PAR PROGRAMME ET PROJET
    # ════════════════════════════════════════════════════════
    doc.add_page_break()
    h1(doc, 'III. ANALYSE PAR PROGRAMME ET PAR PROJET')

    n_rows = 1 + sum(1 + len(p['projets']) for p in s['stats_programmes']) + 1
    tbl_pg = doc.add_table(rows=max(n_rows, 2), cols=6)
    borders(tbl_pg)
    repeat_hdr(tbl_pg.rows[0])
    for i, txt in enumerate(['Code', 'Intitulé', 'Activités', 'Poids (%)',
                              'Budget (F CFA)', 'FP / FADeC / PTFs']):
        hdr_cell(tbl_pg.cell(0, i), txt)
    ri = 1
    for prog in s['stats_programmes']:
        src_detail = f"FP:{_fmt(prog['fp'])} | FADeC:{_fmt(prog['fadec'])} | PTFs:{_fmt(prog['ptfs'])}"
        for ci, (val, al, bd) in enumerate([
            (prog['code'],                               'center', True),
            (f"Programme {prog['num']} : {prog['nom']}", 'left',   True),
            (str(prog['nb_activites']),                  'center', True),
            (_fr(prog['poids']),                          'center', True),
            (_fmt(prog['budget']),                        'right',  True),
            (src_detail,                                  'left',   False),
        ]):
            data_cell(tbl_pg.rows[ri].cells[ci], val, align=al, bold=bd, size=9)
            shade(tbl_pg.rows[ri].cells[ci], 'F4B183')
        ri += 1
        for proj in prog['projets']:
            src_p = f"FP:{_fmt(proj['fp'])} | FADeC:{_fmt(proj['fadec'])} | PTFs:{_fmt(proj['ptfs'])}"
            for ci, (val, al, bd) in enumerate([
                (proj['code'],                                  'center', False),
                (f"  Projet {proj['code']} : {proj['nom']}",  'left',   False),
                (str(proj['nb_activites']),                    'center', False),
                (_fr(proj['poids']),                            'center', False),
                (_fmt(proj['budget']),                          'right',  False),
                (src_p,                                         'left',   False),
            ]):
                data_cell(tbl_pg.rows[ri].cells[ci], val, align=al, bold=bd, size=9)
                shade(tbl_pg.rows[ri].cells[ci], 'FFFF00')
            ri += 1
    for ci, (val, al, bd) in enumerate([
        ('TOTAL', 'center', True), ('', 'left', True),
        (str(s['nb_activites']), 'center', True), ('', 'center', True),
        (_fmt(s['total_budget']), 'right', True), ('', 'left', True),
    ]):
        data_cell(tbl_pg.rows[ri].cells[ci], val, align=al, bold=bd, size=9)
        shade(tbl_pg.rows[ri].cells[ci], 'AED6F1')
    doc.add_paragraph()

    # ════════════════════════════════════════════════════════
    # IV. PAR DIRECTION
    # ════════════════════════════════════════════════════════
    doc.add_page_break()
    h1(doc, 'IV. ACTIVITÉS PAI PAR DIRECTION RESPONSABLE')

    actifs = [d for d in s['stats_directions'] if d['nb_activites'] > 0]
    tbl_dir = doc.add_table(rows=len(actifs) + 2, cols=5)
    borders(tbl_dir)
    repeat_hdr(tbl_dir.rows[0])
    for i, txt in enumerate(['Code', 'Direction', 'Nb activités', 'Budget (F CFA)', 'FP / FADeC / PTFs']):
        hdr_cell(tbl_dir.cell(0, i), txt, bg='27AE60' if i >= 2 else '2C3E50')
    for ri2, d in enumerate(actifs, 1):
        src_d = f"RP:{_fmt(d['rp'])} | FA:{_fmt(d['fa'])} | FNA:{_fmt(d['fn'])} | AP:{_fmt(d['ap'])} | AF:{_fmt(d['af'])}"
        data_cell(tbl_dir.rows[ri2].cells[0], d['code'],            align='center', size=9)
        data_cell(tbl_dir.rows[ri2].cells[1], d['nom'],             align='left',   size=9)
        data_cell(tbl_dir.rows[ri2].cells[2], str(d['nb_activites']), align='center', bold=True, size=9)
        data_cell(tbl_dir.rows[ri2].cells[3], _fmt(d['budget']),    align='right',  size=9)
        data_cell(tbl_dir.rows[ri2].cells[4], src_d,                align='left',   size=9)
        col = 'EBF5FB' if ri2 % 2 == 0 else 'F9F9F9'
        for cell in tbl_dir.rows[ri2].cells:
            shade(cell, col)
    tot_d = tbl_dir.rows[-1]
    data_cell(tot_d.cells[0], '—',  align='center',  bold=True, size=9)
    data_cell(tot_d.cells[1], 'TOTAL', align='left', bold=True, size=9)
    data_cell(tot_d.cells[2], str(s['nb_activites']), align='center', bold=True, size=9)
    data_cell(tot_d.cells[3], _fmt(s['total_budget']), align='right', bold=True, size=9)
    data_cell(tot_d.cells[4], '', align='left', size=9)
    for cell in tot_d.cells:
        shade(cell, 'AED6F1')

    chart_dir = _chart_directions_pai(s['stats_directions'])
    if chart_dir:
        doc.add_picture(chart_dir, width=Cm(14))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # ════════════════════════════════════════════════════════
    # V. LISTE DÉTAILLÉE [paysage]
    # ════════════════════════════════════════════════════════
    if s['liste_activites']:
        switch_landscape(doc)
        h1(doc, "V. LISTE DÉTAILLÉE DES ACTIVITÉS D'INVESTISSEMENT")
        interp(doc,
            f"Légende : FP = Fonds Propres | FADeC = FA + FNA | PTFs = Partenaires Techniques et Financiers. "
            f"Total : {_fmt(s['total_budget'])} F CFA pour {s['nb_activites']} activité(s)."
        )
        cols_det = ['Code PAI', 'Désignation', 'Dir.', 'Période', 'Poids\n(%)',
                    'FP (F CFA)', 'FADeC (F CFA)', 'PTFs (F CFA)', 'Total (F CFA)']
        tbl_det = doc.add_table(rows=len(s['liste_activites']) + 2, cols=len(cols_det))
        borders(tbl_det)
        repeat_hdr(tbl_det.rows[0])
        for ci, txt in enumerate(cols_det):
            hdr_cell(tbl_det.cell(0, ci), txt, bg='1A5276')
        tot_fp = tot_fadec = tot_ptfs = tot_bud = 0
        for ri3, a in enumerate(s['liste_activites'], 1):
            col_bg = 'EBF5FB' if ri3 % 2 == 0 else 'F9F9F9'
            for ci, (val, al) in enumerate([
                (a['code_pai'],      'center'),
                (a['nom'],           'left'),
                (a['direction'],     'center'),
                (a['periode'],       'center'),
                (_fr(a['poids']),    'center'),
                (_fmt(a['fp']),      'right'),
                (_fmt(a['fadec']),   'right'),
                (_fmt(a['ptfs']),    'right'),
                (_fmt(a['budget']),  'right'),
            ]):
                data_cell(tbl_det.rows[ri3].cells[ci], val, align=al, size=8)
                shade(tbl_det.rows[ri3].cells[ci], col_bg)
            tot_fp    += a['fp']
            tot_fadec += a['fadec']
            tot_ptfs  += a['ptfs']
            tot_bud   += a['budget']
        tot_r2 = tbl_det.rows[-1]
        for ci, (val, al) in enumerate([
            ('TOTAL', 'center'), ('', 'left'), ('', 'center'), ('', 'center'), ('', 'center'),
            (_fmt(tot_fp), 'right'), (_fmt(tot_fadec), 'right'),
            (_fmt(tot_ptfs), 'right'), (_fmt(tot_bud), 'right'),
        ]):
            data_cell(tot_r2.cells[ci], val, align=al, bold=True, size=8)
            shade(tot_r2.cells[ci], 'AED6F1')
        switch_portrait(doc)

    # ════════════════════════════════════════════════════════
    # VI. CIBLES TRIMESTRIELLES
    # ════════════════════════════════════════════════════════
    doc.add_page_break()
    h1(doc, "VI. CIBLES TRIMESTRIELLES D'EXÉCUTION")
    interp(doc,
        "Les cibles représentent le taux d'avancement attendu si chaque activité est réalisée "
        "à 100 % dans ses délais planifiés. Calcul : pro-ratisation linéaire mensuelle sur "
        "[période_début → période_fin] de chaque activité, remontée par poids PAI "
        "(activité → projet → programme). "
        "Code couleur — Rouge : < 25 % | Orange : 25–50 % | Bleu : 50–75 % | Vert : ≥ 75 %."
    )
    doc.add_paragraph()

    if cibles:
        h2(doc, 'A. Par Programme et par Projet')
        progs_c  = cibles.get('programmes', [])
        global_c = cibles.get('global', {1: 0, 2: 0, 3: 0, 4: 0})
        n_rows_c = 1 + sum(1 + len(p.get('projets', [])) for p in progs_c)
        tbl_c = doc.add_table(rows=n_rows_c + 1, cols=5)
        borders(tbl_c)
        repeat_hdr(tbl_c.rows[0])
        for i, txt in enumerate(['Niveau', 'T1 (fin Mars)', 'T2 (fin Juin)', 'T3 (fin Sept.)', 'T4 (fin Déc.)']):
            hdr_cell(tbl_c.cell(0, i), txt, bg='1A5276')
        ri_c = 1
        r_gl = tbl_c.rows[ri_c]
        data_cell(r_gl.cells[0], 'PAI Global', bold=True, size=9)
        shade(r_gl.cells[0], 'F4B183')
        for ti, trim in enumerate((1, 2, 3, 4), 1):
            cible_cell(r_gl.cells[ti], global_c.get(trim, 0))
        ri_c += 1
        for prog in progs_c:
            r_p = tbl_c.rows[ri_c]
            data_cell(r_p.cells[0], f"Programme {prog['code']} — {prog['nom']}", bold=True, size=9)
            shade(r_p.cells[0], 'D4E6F1')
            for ti, trim in enumerate((1, 2, 3, 4), 1):
                cible_cell(r_p.cells[ti], prog['cibles'].get(trim, 0))
            ri_c += 1
            for proj in prog.get('projets', []):
                r_pj = tbl_c.rows[ri_c]
                data_cell(r_pj.cells[0], f"   Projet {proj['code']} — {proj['nom']}", size=9)
                shade(r_pj.cells[0], 'EBF5FB')
                for ti, trim in enumerate((1, 2, 3, 4), 1):
                    cible_cell(r_pj.cells[ti], proj['cibles'].get(trim, 0))
                ri_c += 1
        doc.add_paragraph()

        g = global_c
        interp(doc,
            f"Selon la planification, le PAI {annee.annee} devrait atteindre "
            f"{_fr(g.get(1, 0))} % à fin T1, {_fr(g.get(2, 0))} % à fin T2 "
            f"et {_fr(g.get(3, 0))} % à fin T3."
        )

    if cibles_directions:
        h2(doc, 'B. Par Direction responsable')
        tbl_cd = doc.add_table(rows=len(cibles_directions) + 1, cols=7)
        borders(tbl_cd)
        repeat_hdr(tbl_cd.rows[0])
        for i, txt in enumerate(['Code', 'Direction', 'Nb activités',
                                  'T1 (fin Mars)', 'T2 (fin Juin)',
                                  'T3 (fin Sept.)', 'T4 (fin Déc.)']):
            hdr_cell(tbl_cd.cell(0, i), txt, bg='1A5276')
        for ri_d, d in enumerate(cibles_directions, 1):
            r_d = tbl_cd.rows[ri_d]
            data_cell(r_d.cells[0], d['code'],            align='center', bold=True, size=9)
            data_cell(r_d.cells[1], d['nom'],              align='left',             size=9)
            data_cell(r_d.cells[2], str(d['nb_activites']), align='center',          size=9)
            for trim in (1, 2, 3, 4):
                cible_cell(r_d.cells[trim + 2], d['cibles'].get(trim, 0))
            col = 'F4F6F7' if ri_d % 2 == 0 else 'EBF5FB'
            for ci in range(3):
                shade(r_d.cells[ci], col)
        doc.add_paragraph()

    return doc


# ── Export Excel ──────────────────────────────────────────────────────────────

def _build_excel_pai(annee, s):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb   = Workbook()
    thin = Side(style='thin')
    brd  = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctr  = Alignment(horizontal='center', vertical='center', wrap_text=True)
    lft  = Alignment(horizontal='left',   vertical='center', wrap_text=True)

    F_TITRE = PatternFill("solid", fgColor="1F4E79")
    F_HDR   = PatternFill("solid", fgColor="2E75B6")
    F_PROG  = PatternFill("solid", fgColor="F4B183")
    F_PROJ  = PatternFill("solid", fgColor="FFFF00")
    F_TOTAL = PatternFill("solid", fgColor="AED6F1")
    F_DIR   = PatternFill("solid", fgColor="EBF5FB")

    def _titre(ws, texte):
        ws.merge_cells('A1:J1')
        ws['A1'] = texte
        ws['A1'].font = Font(bold=True, size=13, color="FFFFFF")
        ws['A1'].fill = F_TITRE
        ws['A1'].alignment = ctr

    def _entete(ws, cols, row=2, fill=None):
        fill = fill or F_HDR
        for ci, col in enumerate(cols, 1):
            c = ws.cell(row=row, column=ci, value=col)
            c.font = Font(bold=True, color="FFFFFF", size=9)
            c.fill = fill; c.alignment = ctr; c.border = brd
        return row + 1

    def _cell(ws, r, c, v, fill=None, bold=False, align=None):
        cell = ws.cell(row=r, column=c, value=v)
        if fill:  cell.fill = fill
        cell.font      = Font(bold=bold, size=9)
        cell.alignment = align or lft
        cell.border    = brd
        return cell

    def _num(v):
        try: return round(float(v or 0))
        except: return 0

    # ── Feuille 1 : Résumé ─────────────────────────────────────────────────
    ws1 = wb.active; ws1.title = "Résumé général"
    _titre(ws1, f"STATISTIQUES PAI {annee.annee} — Mairie d'Adja-Ouèrè — Résumé")
    rows_resume = [
        ("Budget total PAI",             f"{_fmt(s['total_budget'])} F CFA"),
        ("",                             ""),
        ("Ressources Propres (RP)",      f"{_fmt(s['total_rp'])} F CFA"),
        ("FADeC Affecté (FA)",           f"{_fmt(s['total_fa'])} F CFA"),
        ("FADeC Non Affecté (FNA)",      f"{_fmt(s['total_fn'])} F CFA"),
        ("Appui Partenaires (AP)",       f"{_fmt(s['total_ap'])} F CFA"),
        ("Autres Financements (AF)",     f"{_fmt(s['total_af'])} F CFA"),
        ("",                             ""),
        ("Nombre de Programmes",         s['nb_programmes']),
        ("Nombre de Projets",            s['nb_projets']),
        ("Activités d'investissement",   s['nb_activites']),
    ]
    r = 2
    for lbl, val in rows_resume:
        if lbl == "":
            r += 1; continue
        ws1.cell(row=r, column=1, value=lbl).font = Font(size=9)
        ws1.cell(row=r, column=1).border = brd
        ws1.cell(row=r, column=1).alignment = lft
        c2 = ws1.cell(row=r, column=2, value=val)
        c2.font = Font(size=9, bold=True); c2.border = brd; c2.alignment = ctr
        r += 1
    ws1.column_dimensions['A'].width = 36
    ws1.column_dimensions['B'].width = 26

    # ── Feuille 2 : Par Programme ──────────────────────────────────────────
    ws2 = wb.create_sheet("Par Programme")
    _titre(ws2, f"STATISTIQUES PAR PROGRAMME — PAI {annee.annee}")
    cols2 = ["Code", "Programme / Projet", "Nb Projets", "Nb Activités",
             "Poids (%)", "Budget (F CFA)", "FP (F CFA)", "FADeC (F CFA)", "PTFs (F CFA)"]
    next_r = _entete(ws2, cols2, row=2)
    for pe in s['stats_programmes']:
        _cell(ws2, next_r, 1, pe['code'],           F_PROG, bold=True, align=ctr)
        _cell(ws2, next_r, 2, f"Programme {pe['num']} : {pe['nom']}", F_PROG, bold=True)
        _cell(ws2, next_r, 3, pe['nb_projets'],     F_PROG, bold=True, align=ctr)
        _cell(ws2, next_r, 4, pe['nb_activites'],   F_PROG, bold=True, align=ctr)
        _cell(ws2, next_r, 5, pe['poids'],          F_PROG, bold=True, align=ctr)
        _cell(ws2, next_r, 6, _num(pe['budget']),   F_PROG, bold=True, align=ctr)
        _cell(ws2, next_r, 7, _num(pe['fp']),       F_PROG, bold=True, align=ctr)
        _cell(ws2, next_r, 8, _num(pe['fadec']),    F_PROG, bold=True, align=ctr)
        _cell(ws2, next_r, 9, _num(pe['ptfs']),     F_PROG, bold=True, align=ctr)
        next_r += 1
        for prj in pe.get('projets', []):
            _cell(ws2, next_r, 1, prj['code'],                      F_PROJ, align=ctr)
            _cell(ws2, next_r, 2, f"  Projet {prj['code']} : {prj['nom']}", F_PROJ)
            _cell(ws2, next_r, 3, '',                                F_PROJ, align=ctr)
            _cell(ws2, next_r, 4, prj['nb_activites'],              F_PROJ, align=ctr)
            _cell(ws2, next_r, 5, prj['poids'],                     F_PROJ, align=ctr)
            _cell(ws2, next_r, 6, _num(prj['budget']),              F_PROJ, align=ctr)
            _cell(ws2, next_r, 7, _num(prj['fp']),                  F_PROJ, align=ctr)
            _cell(ws2, next_r, 8, _num(prj['fadec']),               F_PROJ, align=ctr)
            _cell(ws2, next_r, 9, _num(prj['ptfs']),                F_PROJ, align=ctr)
            next_r += 1
    _cell(ws2, next_r, 1, "TOTAL",                 F_TOTAL, bold=True, align=ctr)
    _cell(ws2, next_r, 2, "",                       F_TOTAL)
    _cell(ws2, next_r, 3, "",                       F_TOTAL, align=ctr)
    _cell(ws2, next_r, 4, s['nb_activites'],        F_TOTAL, bold=True, align=ctr)
    _cell(ws2, next_r, 5, "",                       F_TOTAL, align=ctr)
    _cell(ws2, next_r, 6, _num(s['total_budget']),  F_TOTAL, bold=True, align=ctr)
    _cell(ws2, next_r, 7, _num(s['total_fp']),      F_TOTAL, bold=True, align=ctr)
    _cell(ws2, next_r, 8, _num(s['total_fadec']),   F_TOTAL, bold=True, align=ctr)
    _cell(ws2, next_r, 9, _num(s['total_ptfs']),    F_TOTAL, bold=True, align=ctr)
    for i, w in enumerate([10, 50, 12, 14, 10, 20, 18, 18, 16], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = 'A3'

    # ── Feuille 3 : Par Direction ──────────────────────────────────────────
    ws3 = wb.create_sheet("Par Direction")
    _titre(ws3, f"ACTIVITÉS PAI PAR DIRECTION RESPONSABLE — {annee.annee}")
    cols3 = ["Code", "Direction", "Nb Activités", "Budget (F CFA)", "FP (F CFA)",
             "FADeC (F CFA)", "PTFs (F CFA)"]
    next_r = _entete(ws3, cols3, row=2)
    for d in s['stats_directions']:
        bg = F_DIR if d['nb_activites'] > 0 else None
        _cell(ws3, next_r, 1, d['code'],           bg, align=ctr)
        _cell(ws3, next_r, 2, d['nom'],             bg)
        _cell(ws3, next_r, 3, d['nb_activites'],   bg, bold=(d['nb_activites'] > 0), align=ctr)
        _cell(ws3, next_r, 4, _num(d['budget']),   bg, align=ctr)
        _cell(ws3, next_r, 5, _num(d['rp']),       bg, align=ctr)
        _cell(ws3, next_r, 6, _num(d['fadec']),    bg, align=ctr)
        _cell(ws3, next_r, 7, _num(d['ptfs']),     bg, align=ctr)
        next_r += 1
    for i, w in enumerate([10, 45, 14, 18, 16, 16, 14], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = 'A3'

    # ── Feuille 4 : Liste activités ────────────────────────────────────────
    ws4 = wb.create_sheet("Liste activités PAI")
    _titre(ws4, f"LISTE DES ACTIVITÉS D'INVESTISSEMENT — PAI {annee.annee}")
    cols4 = ["Code PAI", "Désignation", "Direction", "Localisation",
             "Période", "Poids (%)", "Indicateurs",
             "RP (F CFA)", "FA (F CFA)", "FNA (F CFA)", "AP (F CFA)", "AF (F CFA)",
             "Budget total (F CFA)"]
    next_r = _entete(ws4, cols4, row=2)
    for a in s['liste_activites']:
        _cell(ws4, next_r,  1, a['code_pai'],     None, align=ctr)
        _cell(ws4, next_r,  2, a['nom'])
        _cell(ws4, next_r,  3, a['direction'],    None, align=ctr)
        _cell(ws4, next_r,  4, a['localisation'])
        _cell(ws4, next_r,  5, a['periode'],      None, align=ctr)
        _cell(ws4, next_r,  6, a['poids'],        None, align=ctr)
        _cell(ws4, next_r,  7, a['indicateurs'])
        _cell(ws4, next_r,  8, _num(a['rp']),     None, align=ctr)
        _cell(ws4, next_r,  9, _num(a['fa']),     None, align=ctr)
        _cell(ws4, next_r, 10, _num(a['fn']),     None, align=ctr)
        _cell(ws4, next_r, 11, _num(a['ap']),     None, align=ctr)
        _cell(ws4, next_r, 12, _num(a['af']),     None, align=ctr)
        _cell(ws4, next_r, 13, _num(a['budget']), None, bold=True, align=ctr)
        next_r += 1
    _cell(ws4, next_r,  1, "TOTAL",                   F_TOTAL, bold=True, align=ctr)
    for ci in (2, 3, 4, 5, 6, 7): _cell(ws4, next_r, ci, '', F_TOTAL, align=ctr)
    _cell(ws4, next_r,  8, _num(s['total_rp']),        F_TOTAL, bold=True, align=ctr)
    _cell(ws4, next_r,  9, _num(s['total_fa']),        F_TOTAL, bold=True, align=ctr)
    _cell(ws4, next_r, 10, _num(s['total_fn']),        F_TOTAL, bold=True, align=ctr)
    _cell(ws4, next_r, 11, _num(s['total_ap']),        F_TOTAL, bold=True, align=ctr)
    _cell(ws4, next_r, 12, _num(s['total_af']),        F_TOTAL, bold=True, align=ctr)
    _cell(ws4, next_r, 13, _num(s['total_budget']),    F_TOTAL, bold=True, align=ctr)
    for i, w in enumerate([10, 40, 12, 22, 14, 9, 28, 14, 14, 14, 14, 14, 18], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    ws4.freeze_panes = 'A3'

    return wb


# ── Routes ────────────────────────────────────────────────────────────────────

@statspai_bp.route('/')
@admin_only
def index():
    import base64
    annee = get_annee()
    if not annee:
        return render_template('statspai/index.html', annee=None)
    s                 = _compute_stats_pai(annee)
    cibles            = _compute_cibles_pai(s['pai_data'])
    cibles_directions = _compute_cibles_directions_pai(s['pai_data'])
    chart_buf = _chart_sources_pai(s['total_rp'], s['total_fa'], s['total_fn'], s['total_ap'], s['total_af'])
    chart_sources_b64 = (base64.b64encode(chart_buf.read()).decode('utf-8')
                         if chart_buf else None)
    return render_template('statspai/index.html',
                           annee=annee, cibles=cibles,
                           cibles_directions=cibles_directions,
                           chart_sources_b64=chart_sources_b64, **s)


@statspai_bp.route('/rapport/word')
@admin_only
def rapport_word():
    annee = get_annee()
    if not annee:
        flash('Aucune année active.', 'danger')
        return redirect(url_for('statspai.index'))
    try:
        s                 = _compute_stats_pai(annee)
        cibles            = _compute_cibles_pai(s['pai_data'])
        cibles_directions = _compute_cibles_directions_pai(s['pai_data'])
        static_img        = os.path.join(current_app.root_path, 'static', 'img')
        doc = _build_word_pai(annee, s, static_img,
                              cibles=cibles, cibles_directions=cibles_directions)
        output = io.BytesIO()
        doc.save(output); output.seek(0)
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name=f"Rapport_Stats_PAI_{annee.annee}.docx",
        )
    except Exception as e:
        flash(f'Erreur lors de la génération du rapport : {e}', 'danger')
        return redirect(url_for('statspai.index'))


@statspai_bp.route('/rapport/excel')
@admin_only
def rapport_excel():
    annee = get_annee()
    if not annee:
        flash('Aucune année active.', 'danger')
        return redirect(url_for('statspai.index'))
    try:
        s  = _compute_stats_pai(annee)
        wb = _build_excel_pai(annee, s)
        output = io.BytesIO()
        wb.save(output); output.seek(0)
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f"Statistiques_PAI_{annee.annee}.xlsx",
        )
    except Exception as e:
        flash(f'Erreur lors de la génération Excel : {e}', 'danger')
        return redirect(url_for('statspai.index'))
