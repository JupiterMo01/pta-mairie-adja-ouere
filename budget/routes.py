import io
from datetime import datetime
from flask import render_template, request, abort, jsonify, send_file
from flask_login import login_required, current_user
from models import db, BudgetExecution
from utils import get_annee
from . import budget_bp

TYPES = ('fonctionnement', 'investissement')
LABELS = {
    'fonctionnement': 'Dépenses de Fonctionnement',
    'investissement': "Dépenses d'Investissement",
}


# ─── Helpers ─────────────────────────────────────────────────────────────────

def _peut_editer():
    if current_user.role == 'admin_editeur':
        return True
    if current_user.role == 'direction' and current_user.direction \
            and current_user.direction.code == 'DAAF':
        return True
    if current_user.role == 'service' and current_user.service \
            and current_user.service.code == 'SBFC':
        return True
    return False


def _taux(num, denom):
    return round(num / denom * 100, 2) if denom else 0.0


def _get_or_create(annee_id, trimestre, type_depense):
    rec = BudgetExecution.query.filter_by(
        annee_id=annee_id, trimestre=trimestre, type_depense=type_depense
    ).first()
    if not rec:
        rec = BudgetExecution(annee_id=annee_id, trimestre=trimestre,
                              type_depense=type_depense)
        db.session.add(rec)
        db.session.flush()
    return rec


def _build_ligne(rec):
    prev  = rec.previsions      or 0.0
    eng   = rec.montant_engage  or 0.0
    mand  = rec.montant_mandate or 0.0
    paye  = rec.montant_paye    or 0.0
    return {
        'record':   rec,
        'previsions': prev,
        'engage':   eng,
        'mandate':  mand,
        'paye':     paye,
        'taux_eng':  _taux(eng,  prev),
        'taux_mand': _taux(mand, eng),
        'taux_pay':  _taux(paye, mand),
        'previsions_verrouillees': bool(rec.previsions_verrouillees),
    }


def _build_data(annee, trimestre):
    lignes = {}
    for t in TYPES:
        rec = _get_or_create(annee.id, trimestre, t)
        lignes[t] = _build_ligne(rec)
    db.session.commit()

    tp  = sum(lignes[t]['previsions'] for t in TYPES)
    te  = sum(lignes[t]['engage']     for t in TYPES)
    tm  = sum(lignes[t]['mandate']    for t in TYPES)
    tpa = sum(lignes[t]['paye']       for t in TYPES)
    lignes['total'] = {
        'previsions': tp,
        'engage':     te,
        'mandate':    tm,
        'paye':       tpa,
        'taux_eng':   _taux(te,  tp),
        'taux_mand':  _taux(tm,  te),
        'taux_pay':   _taux(tpa, tm),
    }
    return lignes


# ─── Routes ──────────────────────────────────────────────────────────────────

@budget_bp.route('/')
@login_required
def index():
    annee = get_annee()
    if not annee:
        abort(404)
    trimestre = request.args.get('t', 1, type=int)
    if trimestre not in (1, 2, 3, 4):
        trimestre = 1
    lignes = _build_data(annee, trimestre)
    return render_template('budget/index.html',
                           annee=annee, trimestre=trimestre, lignes=lignes,
                           peut_editer=_peut_editer(),
                           is_admin=current_user.role == 'admin_editeur',
                           LABELS=LABELS)


# ─── Édition AJAX ────────────────────────────────────────────────────────────

@budget_bp.route('/edit', methods=['POST'])
@login_required
def edit():
    if not _peut_editer():
        abort(403)

    annee = get_annee()
    if not annee:
        abort(404)

    type_depense = request.form.get('type_depense', '')
    if type_depense not in TYPES:
        return jsonify({'ok': False, 'error': 'Type de dépense invalide.'}), 400

    try:
        trimestre = int(request.form.get('trimestre', 0))
        assert trimestre in (1, 2, 3, 4)
    except Exception:
        return jsonify({'ok': False, 'error': 'Trimestre invalide.'}), 400

    try:
        prev_new  = float(request.form.get('previsions', 0) or 0)
        eng_new   = float(request.form.get('montant_engage',  0) or 0)
        mand_new  = float(request.form.get('montant_mandate', 0) or 0)
        paye_new  = float(request.form.get('montant_paye',    0) or 0)
    except ValueError:
        return jsonify({'ok': False, 'error': 'Valeurs numériques invalides.'}), 400

    def _err(msg):
        return jsonify({'ok': False, 'error': msg}), 400

    rec       = _get_or_create(annee.id, trimestre, type_depense)
    is_admin  = current_user.role == 'admin_editeur'

    prev_act  = rec.previsions      or 0.0
    eng_act   = rec.montant_engage  or 0.0
    mand_act  = rec.montant_mandate or 0.0
    paye_act  = rec.montant_paye    or 0.0

    # Prévisions : verrouillées pour non-admin après première saisie
    if not is_admin and rec.previsions_verrouillees:
        prev_new = prev_act

    # Contraintes hiérarchiques
    if eng_new > (prev_new if prev_new else float('inf')):
        return _err("Le montant engagé ne peut pas dépasser les prévisions.")
    if mand_new > eng_new:
        return _err("Le montant mandaté ne peut pas dépasser le montant engagé.")
    if paye_new > mand_new:
        return _err("Le montant payé ne peut pas dépasser le montant mandaté.")

    # Non-admin : montants ne peuvent qu'augmenter
    if not is_admin:
        if eng_new < eng_act:
            return _err(f"Le montant engagé ne peut pas diminuer "
                        f"(valeur actuelle : {eng_act:,.0f} F CFA).")
        if mand_new < mand_act:
            return _err(f"Le montant mandaté ne peut pas diminuer "
                        f"(valeur actuelle : {mand_act:,.0f} F CFA).")
        if paye_new < paye_act:
            return _err(f"Le montant payé ne peut pas diminuer "
                        f"(valeur actuelle : {paye_act:,.0f} F CFA).")

    rec.previsions      = prev_new
    rec.montant_engage  = eng_new
    rec.montant_mandate = mand_new
    rec.montant_paye    = paye_new

    # Verrouiller dès la première saisie non nulle, quel que soit le rôle.
    if prev_new > 0:
        rec.previsions_verrouillees = True

    db.session.commit()

    ligne = _build_ligne(rec)
    return jsonify({
        'ok': True,
        'previsions':   ligne['previsions'],
        'engage':       ligne['engage'],
        'mandate':      ligne['mandate'],
        'paye':         ligne['paye'],
        'taux_eng':     ligne['taux_eng'],
        'taux_mand':    ligne['taux_mand'],
        'taux_pay':     ligne['taux_pay'],
        'previsions_verrouillees': ligne['previsions_verrouillees'],
    })


# ─── Print ───────────────────────────────────────────────────────────────────

@budget_bp.route('/print')
@login_required
def print_view():
    annee = get_annee()
    if not annee:
        abort(404)
    tris = request.args.getlist('t', type=int)
    tris = [t for t in tris if t in (1, 2, 3, 4)] or [1]
    sections = []
    for t in tris:
        lignes = _build_data(annee, t)
        sections.append({'trimestre': t, 'lignes': lignes})
    return render_template('budget/print.html',
                           annee=annee, sections=sections, LABELS=LABELS)


# ─── Export Excel ────────────────────────────────────────────────────────────

@budget_bp.route('/export')
@login_required
def export_excel():
    import openpyxl, os as _os
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from flask import current_app

    annee = get_annee()
    if not annee:
        abort(404)

    tris = request.args.getlist('t', type=int)
    tris = [t for t in tris if t in (1, 2, 3, 4)] or [1]

    thin  = Side(style='thin'); med = Side(style='medium')
    brd   = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctr   = Alignment(horizontal='center', vertical='center', wrap_text=True)
    lft   = Alignment(horizontal='left',   vertical='center', wrap_text=True)
    rgt   = Alignment(horizontal='right',  vertical='center', wrap_text=True)

    f_hdr  = PatternFill('solid', fgColor='1F4E79')
    f_fct  = PatternFill('solid', fgColor='D6EAF8')
    f_inv  = PatternFill('solid', fgColor='D5F5E3')
    f_tot  = PatternFill('solid', fgColor='1F4E79')
    f_tit  = PatternFill('solid', fgColor='AED6F1')

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used = set()

    def _sn(name):
        base = name[:31]; n = base; i = 2
        while n in used: n = (base[:28] + f'_{i}')[:31]; i += 1
        used.add(n); return n

    _simg = _os.path.join(current_app.root_path, 'static', 'img')

    def _fill_sheet(ws, tri, lignes):
        # ── En-tête institutionnel ────────────────────────────────────────────
        ws.row_dimensions[1].height = 22
        ws.row_dimensions[2].height = 8
        ws.row_dimensions[3].height = 22
        ws.merge_cells('A1:C3'); ws.merge_cells('D1:F3'); ws.merge_cells('G1:H3')
        c = ws['D1']
        c.value = "BP 02 Adja-Ouèrè\nTél : +229 01 61 91 96 12\nEmail : contact.adjaouere@mairie.bj"
        c.font = Font(bold=True, size=9)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        try:
            from openpyxl.drawing.image import Image as _XI
            from PIL import Image as _PI
            import io as _io2
            _H = 52
            for _path, _anc in [(_os.path.join(_simg, 'bandeau.png'), 'A1'),
                                 (_os.path.join(_simg, 'logo_commune.png'), None)]:
                if not _os.path.exists(_path): continue
                _buf = _io2.BytesIO()
                with _PI.open(_path) as _p: _ow, _oh = _p.size; _p.save(_buf, 'PNG')
                _buf.seek(0); _xi = _XI(_buf); _xi.height = _H; _xi.width = int(_ow*_H/_oh)
                if _anc: ws.add_image(_xi, _anc)
                else:
                    try:
                        from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
                        from openpyxl.drawing.xdr import XDRPositiveSize2D
                        _a = OneCellAnchor()
                        _a._from = AnchorMarker(col=7, colOff=-int(_xi.width*9525), row=0, rowOff=0)
                        _a.ext = XDRPositiveSize2D(int(_xi.width*9525), int(_H*9525))
                        _xi.anchor = _a; ws.add_image(_xi)
                    except Exception: ws.add_image(_xi, 'G1')
        except Exception: pass

        # ── Titre ─────────────────────────────────────────────────────────────
        ws.merge_cells('A4:H4')
        ws['A4'].value = (f"NIVEAU D'EXÉCUTION DES DÉPENSES BUDGÉTAIRES"
                          f" — Exercice {annee.annee} — Trimestre {tri}")
        ws['A4'].font = Font(bold=True, size=12)
        ws['A4'].alignment = Alignment(horizontal='center', vertical='center')
        ws['A4'].fill = f_tit
        for _c in range(1, 9):
            ws.cell(row=4, column=_c).border = Border(
                left=med if _c==1 else Side(style=None),
                right=med if _c==8 else Side(style=None),
                top=med, bottom=med)
        ws.row_dimensions[4].height = 28

        # ── En-têtes tableau ──────────────────────────────────────────────────
        hdrs = ['Libellés', 'Prévisions', 'Engagements', 'Taux (%)',
                'Mandatements', 'Taux (%)', 'Paiement', 'Taux (%)']
        for ci, h in enumerate(hdrs, 1):
            c = ws.cell(row=5, column=ci, value=h)
            c.fill = f_hdr; c.font = Font(bold=True, size=9, color='FFFFFF')
            c.alignment = ctr if ci > 1 else lft; c.border = brd
        ws.row_dimensions[5].height = 20

        # ── Données ───────────────────────────────────────────────────────────
        fills = {'fonctionnement': f_fct, 'investissement': f_inv}
        row = 6
        for t in TYPES:
            lg = lignes[t]
            vals = [LABELS[t], lg['previsions'], lg['engage'], lg['taux_eng'],
                    lg['mandate'], lg['taux_mand'], lg['paye'], lg['taux_pay']]
            for ci, v in enumerate(vals, 1):
                c = ws.cell(row=row, column=ci, value=round(v, 2) if isinstance(v, float) else v)
                c.fill = fills[t]; c.font = Font(bold=True, size=9); c.border = brd
                c.alignment = lft if ci == 1 else ctr
                if ci in (2, 3, 5, 7): c.number_format = '#,##0'
                if ci in (4, 6, 8):    c.number_format = '0.00"%"'
            row += 1

        # Total
        tl = lignes['total']
        tot_vals = ['TOTAL BUDGET', tl['previsions'], tl['engage'], tl['taux_eng'],
                    tl['mandate'], tl['taux_mand'], tl['paye'], tl['taux_pay']]
        for ci, v in enumerate(tot_vals, 1):
            c = ws.cell(row=row, column=ci, value=round(v, 2) if isinstance(v, float) else v)
            c.fill = f_tot; c.font = Font(bold=True, size=9, color='FFFFFF'); c.border = brd
            c.alignment = lft if ci == 1 else ctr
            if ci in (2, 3, 5, 7): c.number_format = '#,##0'
            if ci in (4, 6, 8):    c.number_format = '0.00"%"'
        row += 1

        # Pied
        _ds = datetime.now().strftime('%d/%m/%Y à %H:%M')
        ws.merge_cells(f'A{row}:D{row}')
        c = ws.cell(row=row, column=1, value=f'Exporté le {_ds}')
        c.font = Font(bold=True, size=8); c.alignment = lft
        ws.merge_cells(f'E{row}:H{row}')
        c = ws.cell(row=row, column=5,
                    value='Direction du Développement Local et de la Planification (DDLP)')
        c.font = Font(bold=True, size=8); c.alignment = rgt

        # Largeurs
        for ci, w in enumerate([38, 18, 18, 10, 18, 10, 18, 10], 1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = 'A6'

    for tri in tris:
        lignes = _build_data(annee, tri)
        ws = wb.create_sheet(_sn(f'T{tri}'))
        _fill_sheet(ws, tri, lignes)

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    tris_lbl = '_'.join(f'T{t}' for t in tris)
    fname = f"Budget_{annee.annee}_{tris_lbl}.xlsx"
    return send_file(buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True, download_name=fname)
