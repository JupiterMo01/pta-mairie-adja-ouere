import secrets
import string
from functools import wraps
from markupsafe import Markup, escape
from flask import render_template, redirect, url_for, flash, request, session
from flask_login import login_required, current_user
import json
from models import db, User, Direction, Service, Annee, StructureExterne, PTABackup, Programme, Projet, Activite, Tache, SuiviTache, PeiActivite
from admin import admin_bp
from extensions import limiter
from utils import (log_audit, valider_mdp, requete_pta, LOGIN_PRINCIPAL,
                   est_admin_principal as _est_principal, affectation_financiere)


def _generer_mdp_temp():
    """Génère un mot de passe temporaire sécurisé (10 car. : maj + min + chiffre)."""
    alphabet = string.ascii_letters + string.digits
    while True:
        mdp = ''.join(secrets.choice(alphabet) for _ in range(10))
        if (any(c.isupper() for c in mdp) and
                any(c.islower() for c in mdp) and
                any(c.isdigit() for c in mdp)):
            return mdp


def editeur_required(f):
    """Décorateur unique pour tout accès au panneau admin.
    Réservé au rôle 'admin_editeur' (lecture ET écriture).
    admin_lecteur n'a pas accès à /admin/ — il peut consulter /dashboard/ et les PTA.
    """
    @wraps(f)
    @login_required
    def decorated(*args, **kwargs):
        if current_user.role != 'admin_editeur':
            flash('Accès refusé. Réservé aux administrateurs éditeurs.', 'danger')
            return redirect(url_for('pta.global_pta'))
        return f(*args, **kwargs)
    return decorated


# Alias conservé pour compatibilité interne — même comportement qu'editeur_required
admin_required = editeur_required


def principal_required(f):
    """Actions réservées à l'administrateur principal (maintenance, années, purges, restaurations,
    directions et services, suppression ou désactivation de comptes, mots de passe)."""
    @wraps(f)
    @editeur_required
    def decorated(*args, **kwargs):
        if not _est_principal(current_user):
            log_audit('action_refusee', f"Action réservée à l'administrateur principal "
                                        f"({request.endpoint}) refusée à {current_user.login}")
            flash("Cette action est réservée à l'administrateur principal.", 'danger')
            return redirect(url_for('admin.index'))
        return f(*args, **kwargs)
    return decorated


# ─── Sauvegarde manuelle ────────────────────────────────────────────────────

@admin_bp.route('/backup-now', methods=['POST'])
@limiter.limit('10 per minute')
@editeur_required
def backup_now():
    """Déclenche une sauvegarde immédiate par email (admin_editeur uniquement)."""
    import importlib.util, os
    script = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'backup_pta.py')
    try:
        spec = importlib.util.spec_from_file_location('backup_pta', script)
        mod  = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        mod.run()
        log_audit('sauvegarde', 'Sauvegarde manuelle déclenchée depuis Administration')
        flash('Sauvegarde envoyée par email avec succès.', 'success')
    except FileNotFoundError:
        flash('Script de sauvegarde introuvable. Contactez l\'administrateur.', 'danger')
    except Exception as e:
        flash(f'Erreur lors de la sauvegarde : {e}', 'danger')
    return redirect(url_for('admin.index'))


# ─── Tableau de bord ────────────────────────────────────────────────────────

@admin_bp.route('/')
@admin_required
def index():
    stats = {
        'users': User.query.count(),
        'directions': Direction.query.count(),
        'services': Service.query.count(),
        'annees': Annee.query.count(),
    }
    par_dir, par_svc = _emails_par_entite()
    return render_template('admin/index.html', stats=stats,
                           rappel_directions=Direction.query.order_by(Direction.code).all(),
                           mails_dir=par_dir, mails_svc=par_svc)


# ─── Utilisateurs ───────────────────────────────────────────────────────────

ROLES_NON_ADMIN = ('direction', 'service')


def _refus_principal(user):
    """Le compte de l'administrateur principal n'est modifiable que par lui-même.
    Retourne une redirection si l'action doit être refusée, sinon None."""
    if _est_principal(user) and not _est_principal(current_user):
        log_audit('action_refusee',
                  f"Tentative d'action sur le compte administrateur principal par {current_user.login}")
        flash("Le compte de l'administrateur principal ne peut être modifié que par lui-même.", 'danger')
        return redirect(url_for('admin.users'))
    return None


def _refus_hors_principal(role, direction_id, service_id, actuel=None):
    """Hors administrateur principal : seuls des comptes direction ou service, jamais rattachés à la DAAF
    ni au SBFC (saisie financière du PAI). Retourne une redirection si refusé, sinon None."""
    if _est_principal(current_user):
        return None
    motif = None
    if role not in ROLES_NON_ADMIN:
        motif = "Seul l'administrateur principal peut attribuer un profil administrateur."
    elif affectation_financiere(role, direction_id, service_id) and not (
            actuel and affectation_financiere(actuel.role, actuel.direction_id, actuel.service_id)
            and (actuel.role, actuel.direction_id, actuel.service_id) == (
                role, int(direction_id) if direction_id else None, int(service_id) if service_id else None)):
        motif = ("Seul l'administrateur principal peut rattacher un compte à la DAAF ou au SBFC, "
                 "qui saisissent l'exécution financière du PAI.")
    if motif:
        log_audit('action_refusee', f"{motif} (tentative de {current_user.login})")
        flash(motif, 'danger')
        return redirect(url_for('admin.users'))
    return None


@admin_bp.route('/users')
@admin_required
def users():
    users = User.query.order_by(User.nom).all()
    directions = Direction.query.order_by(Direction.nom).all()
    services = Service.query.order_by(Service.nom).all()
    return render_template('admin/users.html', users=users,
                           directions=directions, services=services,
                           login_principal=LOGIN_PRINCIPAL,
                           suis_principal=_est_principal(current_user))


@admin_bp.route('/users/add', methods=['POST'])
@editeur_required
def user_add():
    nom = request.form.get('nom', '').strip()
    prenom = request.form.get('prenom', '').strip()
    login_val = request.form.get('login', '').strip()
    password = request.form.get('password', '').strip()
    role = request.form.get('role', '').strip()
    direction_id = request.form.get('direction_id') or None
    service_id = request.form.get('service_id') or None

    if not all([nom, prenom, login_val, role]):
        flash('Veuillez remplir tous les champs obligatoires (prénom, nom, identifiant, rôle).', 'danger')
        return redirect(url_for('admin.users'))
    refus = _refus_hors_principal(role, direction_id, service_id)
    if refus:
        return refus

    # Mot de passe : si vide → auto-génération ; sinon validation
    mdp_auto = False
    if not password:
        password = _generer_mdp_temp()
        mdp_auto = True
    else:
        erreur_mdp = valider_mdp(password)
        if erreur_mdp:
            flash(erreur_mdp, 'danger')
            return redirect(url_for('admin.users'))

    if User.query.filter_by(login=login_val).first():
        flash(f"L'identifiant « {login_val} » est déjà utilisé.", 'danger')
        return redirect(url_for('admin.users'))

    import re as _re
    email_raw = request.form.get('email', '').strip()
    if email_raw and not _re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', email_raw):
        flash("Format d'adresse email invalide.", 'danger')
        return redirect(url_for('admin.users'))
    email_val = email_raw or None

    user = User(
        nom=nom, prenom=prenom, login=login_val, role=role,
        direction_id=int(direction_id) if direction_id else None,
        service_id=int(service_id) if service_id else None,
        email=email_val,
    )
    user.set_password(password)
    db.session.add(user)
    db.session.commit()
    log_audit('user_cree', f"Compte créé : {prenom} {nom} ({role})")
    lbl_mdp = 'Mot de passe généré automatiquement' if mdp_auto else 'Mot de passe défini'
    flash(Markup(
        f'Utilisateur <strong>{escape(prenom)} {escape(nom)}</strong> créé.<br>'
        f'{lbl_mdp} : '
        f'<code class="fs-6 fw-bold px-2 py-1 bg-light border rounded">{escape(password)}</code><br>'
        f'<span class="text-muted small">Communiquez-le à l\'utilisateur. '
        f'Il pourra le modifier via son menu en haut à droite.</span>'
    ), 'success')
    return redirect(url_for('admin.users'))


@admin_bp.route('/users/<int:user_id>/edit', methods=['GET', 'POST'])
@editeur_required
def user_edit(user_id):
    user = db.get_or_404(User, user_id)
    refus = _refus_principal(user)
    if refus:
        return refus
    principal = _est_principal(current_user)
    if not principal and user.role not in ROLES_NON_ADMIN:
        log_audit('action_refusee', f"Modification du compte administrateur {user.login} refusée à {current_user.login}")
        flash("Seul l'administrateur principal peut modifier un compte administrateur.", 'danger')
        return redirect(url_for('admin.users'))
    directions = Direction.query.order_by(Direction.nom).all()
    services = Service.query.order_by(Service.nom).all()

    if request.method == 'POST':
        nom = request.form.get('nom', '').strip()
        prenom = request.form.get('prenom', '').strip()
        role = request.form.get('role', '').strip()
        if not all([nom, prenom, role]):
            flash('Nom, prénom et rôle sont obligatoires.', 'danger')
            return redirect(url_for('admin.user_edit', user_id=user_id))
        refus = _refus_hors_principal(role, request.form.get('direction_id') or None,
                                      request.form.get('service_id') or None, actuel=user)
        if refus:
            return refus
        if not principal and request.form.get('password', '').strip():
            flash("Seul l'administrateur principal peut modifier un mot de passe.", 'danger')
            return redirect(url_for('admin.user_edit', user_id=user_id))
        # Empêcher un admin de rétrograder son propre rôle (perte d'accès)
        if user.id == current_user.id and role != 'admin_editeur':
            flash('Vous ne pouvez pas modifier votre propre rôle d\'administrateur.', 'danger')
            return redirect(url_for('admin.user_edit', user_id=user_id))
        user.nom = nom
        user.prenom = prenom
        user.role = role
        direction_id = request.form.get('direction_id') or None
        service_id = request.form.get('service_id') or None
        try:
            user.direction_id = int(direction_id) if direction_id else None
            user.service_id = int(service_id) if service_id else None
        except (ValueError, TypeError):
            flash('Identifiant de direction ou service invalide.', 'danger')
            return redirect(url_for('admin.user_edit', user_id=user_id))
        # Personne ne peut désactiver son propre compte (ni le compte principal) depuis ce formulaire
        # Activer ou désactiver un compte relève de l'administrateur principal
        if principal:
            user.actif = True if (user.id == current_user.id or _est_principal(user)) else ('actif' in request.form)
        import re as _re2
        email_edit = request.form.get('email', '').strip()
        if email_edit and not _re2.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', email_edit):
            flash("Format d'adresse email invalide.", 'danger')
            return redirect(url_for('admin.user_edit', user_id=user_id))
        user.email = email_edit or None
        new_pw = request.form.get('password', '').strip()
        if new_pw:
            erreur_mdp = valider_mdp(new_pw)
            if erreur_mdp:
                flash(erreur_mdp, 'danger')
                return redirect(url_for('admin.user_edit', user_id=user_id))
            user.set_password(new_pw)
        db.session.commit()
        log_audit('user_modifie', f"Utilisateur modifié : {user.prenom} {user.nom} — rôle : {user.role}")
        flash('Utilisateur mis à jour.', 'success')
        return redirect(url_for('admin.users'))

    return render_template('admin/user_form.html', user=user,
                           directions=directions, services=services, suis_principal=principal)


@admin_bp.route('/users/<int:user_id>/toggle', methods=['POST'])
@principal_required
def user_toggle(user_id):
    user = db.get_or_404(User, user_id)
    refus = _refus_principal(user)
    if refus:
        return refus
    if (_est_principal(user) or user.id == current_user.id) and user.actif:
        flash('Impossible de désactiver ce compte (administrateur principal ou votre propre compte).', 'danger')
        return redirect(url_for('admin.users'))
    user.actif = not user.actif
    db.session.commit()
    etat = 'activé' if user.actif else 'désactivé'
    action = 'user_active' if user.actif else 'user_desactive'
    log_audit(action, f"Compte {etat} : {user.prenom} {user.nom}")
    flash(f'{user.prenom} {user.nom} {etat}.', 'success')
    return redirect(url_for('admin.users'))


@admin_bp.route('/users/<int:user_id>/delete', methods=['POST'])
@principal_required
def user_delete(user_id):
    user = db.get_or_404(User, user_id)
    refus = _refus_principal(user)
    if refus:
        return refus
    if _est_principal(user):
        flash('Impossible de supprimer le compte administrateur principal.', 'danger')
        return redirect(url_for('admin.users'))
    if user.id == current_user.id:
        flash('Vous ne pouvez pas supprimer votre propre compte.', 'danger')
        return redirect(url_for('admin.users'))
    nom = f'{user.prenom} {user.nom}'
    role = user.role
    db.session.delete(user)
    db.session.commit()
    log_audit('user_supprime', f"Compte supprimé : {nom} ({role})")
    flash(f'Utilisateur {nom} supprimé.', 'success')
    return redirect(url_for('admin.users'))


@admin_bp.route('/users/<int:user_id>/reset-password', methods=['POST'])
@principal_required
def user_reset_password(user_id):
    user = db.get_or_404(User, user_id)
    refus = _refus_principal(user)
    if refus:
        return refus
    # Génère un mot de passe temporaire aléatoire — l'admin le note et le communique
    mdp_temp = _generer_mdp_temp()
    user.set_password(mdp_temp)
    db.session.commit()
    log_audit('mdp_reinit', f"Mot de passe réinitialisé pour : {user.prenom} {user.nom}")
    # Markup pour afficher le mot de passe en gras sans échappement HTML
    flash(Markup(
        f'Mot de passe de <strong>{escape(user.prenom)} {escape(user.nom)}</strong> réinitialisé.<br>'
        f'Mot de passe temporaire : '
        f'<code class="fs-6 fw-bold px-2 py-1 bg-light border rounded">{escape(mdp_temp)}</code><br>'
        f'<span class="text-muted small">Notez-le et communiquez-le à l\'utilisateur. '
        f'Il pourra le modifier via son menu en haut à droite.</span>'
    ), 'success')
    return redirect(url_for('admin.users'))


# ─── Directions ─────────────────────────────────────────────────────────────

@admin_bp.route('/directions')
@admin_required
def directions():
    directions = Direction.query.order_by(Direction.code).all()
    return render_template('admin/directions.html', directions=directions)


@admin_bp.route('/directions/add', methods=['POST'])
@principal_required
def direction_add():
    code = request.form.get('code', '').strip().upper()
    nom = request.form.get('nom', '').strip()
    if not code or not nom:
        flash('Code et nom obligatoires.', 'danger')
        return redirect(url_for('admin.directions'))
    if Direction.query.filter_by(code=code).first():
        flash(f"Le code « {code} » existe déjà.", 'danger')
        return redirect(url_for('admin.directions'))
    db.session.add(Direction(code=code, nom=nom))
    db.session.commit()
    flash(f'Direction {code} créée.', 'success')
    return redirect(url_for('admin.directions'))


@admin_bp.route('/directions/<int:dir_id>/edit', methods=['POST'])
@principal_required
def direction_edit(dir_id):
    d = db.get_or_404(Direction, dir_id)
    code = request.form.get('code', '').strip().upper()
    nom = request.form.get('nom', '').strip()
    if not code or not nom:
        flash('Code et nom sont obligatoires.', 'danger')
        return redirect(url_for('admin.directions'))
    d.code = code
    d.nom = nom
    db.session.commit()
    flash('Direction mise à jour.', 'success')
    return redirect(url_for('admin.directions'))


@admin_bp.route('/directions/<int:dir_id>/delete', methods=['POST'])
@principal_required
def direction_delete(dir_id):
    d = db.get_or_404(Direction, dir_id)
    if d.services:
        flash('Impossible : des services sont rattachés à cette direction.', 'danger')
        return redirect(url_for('admin.directions'))
    if User.query.filter_by(direction_id=dir_id).first():
        flash('Impossible : des utilisateurs sont rattachés à cette direction.', 'danger')
        return redirect(url_for('admin.directions'))
    db.session.delete(d)
    db.session.commit()
    flash('Direction supprimée.', 'success')
    return redirect(url_for('admin.directions'))


# ─── Services ────────────────────────────────────────────────────────────────

@admin_bp.route('/services')
@admin_required
def services():
    services = Service.query.order_by(Service.code).all()
    directions = Direction.query.order_by(Direction.nom).all()
    return render_template('admin/services.html', services=services, directions=directions)


@admin_bp.route('/services/add', methods=['POST'])
@principal_required
def service_add():
    code = request.form.get('code', '').strip().upper()
    nom = request.form.get('nom', '').strip()
    direction_id = request.form.get('direction_id')
    if not all([code, nom, direction_id]):
        flash('Tous les champs sont obligatoires.', 'danger')
        return redirect(url_for('admin.services'))
    if Service.query.filter_by(code=code).first():
        flash(f"Le code « {code} » existe déjà.", 'danger')
        return redirect(url_for('admin.services'))
    db.session.add(Service(code=code, nom=nom, direction_id=int(direction_id)))
    db.session.commit()
    flash(f'Service {code} créé.', 'success')
    return redirect(url_for('admin.services'))


@admin_bp.route('/services/<int:svc_id>/edit', methods=['POST'])
@principal_required
def service_edit(svc_id):
    s = db.get_or_404(Service, svc_id)
    code = request.form.get('code', '').strip().upper()
    nom = request.form.get('nom', '').strip()
    if not code or not nom:
        flash('Code et nom sont obligatoires.', 'danger')
        return redirect(url_for('admin.services'))
    try:
        direction_id = int(request.form.get('direction_id', 0))
    except (ValueError, TypeError):
        flash('Direction invalide.', 'danger')
        return redirect(url_for('admin.services'))
    s.code = code
    s.nom = nom
    s.direction_id = direction_id
    db.session.commit()
    flash('Service mis à jour.', 'success')
    return redirect(url_for('admin.services'))


@admin_bp.route('/services/<int:svc_id>/delete', methods=['POST'])
@principal_required
def service_delete(svc_id):
    s = db.get_or_404(Service, svc_id)
    if User.query.filter_by(service_id=svc_id).first():
        flash('Impossible : des utilisateurs sont rattachés à ce service.', 'danger')
        return redirect(url_for('admin.services'))
    db.session.delete(s)
    db.session.commit()
    flash('Service supprimé.', 'success')
    return redirect(url_for('admin.services'))


# ─── Années ──────────────────────────────────────────────────────────────────

@admin_bp.route('/annees')
@admin_required
def annees():
    annees = Annee.query.order_by(Annee.annee.desc()).all()
    return render_template('admin/annees.html', annees=annees)


@admin_bp.route('/annees/add', methods=['POST'])
@principal_required
def annee_add():
    val = request.form.get('annee')
    if not val:
        flash('Année obligatoire.', 'danger')
        return redirect(url_for('admin.annees'))
    try:
        val = int(val)
    except (ValueError, TypeError):
        flash("L'année doit être un nombre entier valide (ex : 2026).", 'danger')
        return redirect(url_for('admin.annees'))
    if val < 2000 or val > 2100:
        flash("L'année doit être comprise entre 2000 et 2100.", 'danger')
        return redirect(url_for('admin.annees'))
    if Annee.query.filter_by(annee=val).first():
        flash(f"L'année {val} existe déjà.", 'danger')
        return redirect(url_for('admin.annees'))
    nouvelle = Annee(annee=val, actif=False)
    db.session.add(nouvelle)
    db.session.commit()
    log_audit('annee_cree', f"Année PTA créée : {val}")
    # Proposer immédiatement la copie du PTA de l'année source
    return redirect(url_for('admin.annee_confirm_copie', ann_id=nouvelle.id))


@admin_bp.route('/annees/<int:ann_id>/confirmer-copie')
@principal_required
def annee_confirm_copie(ann_id):
    """Page de confirmation : copier ou non le PTA de l'année source vers la nouvelle année."""
    nouvelle = db.get_or_404(Annee, ann_id)
    # Source = année active, sinon la plus récente différente de la nouvelle
    source = Annee.query.filter_by(actif=True).first()
    if not source or source.id == ann_id:
        source = Annee.query.filter(Annee.id != ann_id)\
                             .order_by(Annee.annee.desc()).first()
    from models import Programme
    nb_prog_source = Programme.query.filter_by(
        annee_id=source.id).count() if source else 0
    return render_template('admin/annee_copie_confirm.html',
                           nouvelle=nouvelle,
                           source=source,
                           nb_prog_source=nb_prog_source)


@admin_bp.route('/annees/<int:ann_id>/copier-pta', methods=['POST'])
@principal_required
def annee_copier_pta(ann_id):
    """Copie la structure complète du PTA (sans suivi) depuis l'année source."""
    nouvelle = db.get_or_404(Annee, ann_id)
    source_id = request.form.get('source_id', type=int)
    if not source_id:
        flash('Année source introuvable.', 'danger')
        return redirect(url_for('admin.annees'))

    source = db.get_or_404(Annee, source_id)
    from models import Programme, Projet, Activite, Tache

    # Vérifier que la nouvelle année est bien vide
    if Programme.query.filter_by(annee_id=ann_id).first():
        flash(f'Le PTA {nouvelle.annee} contient déjà des données. Copie annulée.', 'danger')
        return redirect(url_for('admin.annees'))

    nb_prog = nb_proj = nb_act = nb_tache = 0

    for prog_src in requete_pta(source_id).order_by(Programme.numero).all():
        prog_new = Programme(
            annee_id            = ann_id,
            numero              = prog_src.numero,
            nom                 = prog_src.nom,
            description         = prog_src.description,
            objectif_specifique = prog_src.objectif_specifique,
            poids               = prog_src.poids,
            observations        = prog_src.observations,
        )
        db.session.add(prog_new)
        db.session.flush()   # obtenir prog_new.id
        nb_prog += 1

        for proj_src in sorted(prog_src.projets, key=lambda p: p.numero):
            proj_new = Projet(
                programme_id = prog_new.id,
                numero       = proj_src.numero,
                nom          = proj_src.nom,
                description  = proj_src.description,
                poids        = proj_src.poids,
                observations = proj_src.observations,
            )
            db.session.add(proj_new)
            db.session.flush()
            nb_proj += 1

            for act_src in sorted(proj_src.activites, key=lambda a: a.numero):
                act_new = Activite(
                    projet_id               = proj_new.id,
                    numero                  = act_src.numero,
                    nom                     = act_src.nom,
                    description             = act_src.description,
                    direction_responsable_id= act_src.direction_responsable_id,
                    imputation_budgetaire   = act_src.imputation_budgetaire,
                    ressources_propres      = act_src.ressources_propres,
                    fadec_affecte           = act_src.fadec_affecte,
                    fadec_non_affecte       = act_src.fadec_non_affecte,
                    autres_partenaires      = act_src.autres_partenaires,
                    autres_fonds            = act_src.autres_fonds,
                    details_financement     = act_src.details_financement,
                    acteurs_externes        = act_src.acteurs_externes,
                    periode_debut           = act_src.periode_debut,
                    periode_fin             = act_src.periode_fin,
                    mode_execution          = act_src.mode_execution,
                    type_activite           = act_src.type_activite,
                    poids                   = act_src.poids,
                    observations            = act_src.observations,
                )
                # Many-to-many : services, directions et structures externes intervenants
                act_new.services_intervenants = list(act_src.services_intervenants)
                act_new.directions_associees  = list(act_src.directions_associees)
                act_new.structures_externes   = list(act_src.structures_externes)
                db.session.add(act_new)
                db.session.flush()
                nb_act += 1

                for tache_src in sorted(act_src.taches,
                                        key=lambda t: (t.ordre, t.numero)):
                    tache_new = Tache(
                        activite_id             = act_new.id,
                        numero                  = tache_src.numero,
                        ordre                   = tache_src.ordre,
                        nom                     = tache_src.nom,
                        description             = tache_src.description,
                        poids                   = tache_src.poids,
                        service_responsable_id  = tache_src.service_responsable_id,
                        direction_responsable_id= tache_src.direction_responsable_id,
                        imputation_budgetaire   = tache_src.imputation_budgetaire,
                        ressources_propres      = tache_src.ressources_propres,
                        fadec_affecte           = tache_src.fadec_affecte,
                        fadec_non_affecte       = tache_src.fadec_non_affecte,
                        autres_partenaires      = tache_src.autres_partenaires,
                        autres_fonds            = tache_src.autres_fonds,
                        details_financement     = tache_src.details_financement,
                        acteurs_externes        = tache_src.acteurs_externes,
                        mode_execution          = tache_src.mode_execution,
                        periode_debut           = tache_src.periode_debut,
                        periode_fin             = tache_src.periode_fin,
                        observations            = tache_src.observations,
                    )
                    # Many-to-many : services, directions et structures externes concernés
                    tache_new.services_concernes   = list(tache_src.services_concernes)
                    tache_new.directions_associees = list(tache_src.directions_associees)
                    tache_new.structures_externes  = list(tache_src.structures_externes)
                    db.session.add(tache_new)
                    nb_tache += 1

    # Copier aussi l'objectif général de l'année source
    nouvelle.objectif_general = source.objectif_general
    db.session.commit()
    log_audit('annee_pta_copie', f"PTA {source.annee} → {nouvelle.annee} : {nb_prog} prog., {nb_proj} proj., {nb_act} act., {nb_tache} tâches")

    flash(
        f'Structure du PTA {source.annee} copiée vers {nouvelle.annee} avec succès : '
        f'{nb_prog} programme(s), {nb_proj} projet(s), '
        f'{nb_act} activité(s), {nb_tache} tâche(s). '
        "Aucune donnée de suivi n'a été copiée.",
        'success'
    )
    return redirect(url_for('admin.annees'))


@admin_bp.route('/annees/<int:ann_id>/activate', methods=['POST'])
@principal_required
def annee_activate(ann_id):
    Annee.query.update({'actif': False})
    a = db.get_or_404(Annee, ann_id)
    a.actif = True
    db.session.commit()
    # Mettre à jour la session pour que le badge en haut change immédiatement
    session['annee_id'] = a.id
    session['annee'] = a.annee
    log_audit('annee_activee', f"Année PTA activée : {a.annee}")
    flash(f'Année {a.annee} activée.', 'success')
    return redirect(url_for('admin.annees'))


@admin_bp.route('/annees/<int:ann_id>/deactivate', methods=['POST'])
@principal_required
def annee_deactivate(ann_id):
    a = db.get_or_404(Annee, ann_id)
    a.actif = False
    db.session.commit()
    session.pop('annee_id', None)
    session.pop('annee', None)
    log_audit('annee_desactivee', f"Année PTA désactivée : {a.annee}")
    flash(f'Année {a.annee} désactivée.', 'warning')
    return redirect(url_for('admin.annees'))


@admin_bp.route('/annees/<int:ann_id>/purge-suivi', methods=['POST'])
@limiter.limit('5 per minute')
@principal_required
def annee_purge_suivi(ann_id):
    """Vide uniquement les données de suivi (suivi_taches) pour une année."""
    a = db.get_or_404(Annee, ann_id)
    from models import SuiviTache
    nb = SuiviTache.query.filter_by(annee_id=ann_id).delete()
    db.session.commit()
    log_audit('annee_purge_suivi', f"Suivi PTA {a.annee} purgé : {nb} enregistrement(s) supprimé(s)")
    flash(f'Suivi de {a.annee} vidé : {nb} enregistrement(s) supprimé(s). '
          f'La structure du PTA est conservée.', 'success')
    return redirect(url_for('admin.annees'))


@admin_bp.route('/annees/<int:ann_id>/purge-pta', methods=['POST'])
@limiter.limit('3 per minute')
@principal_required
def annee_purge_pta(ann_id):
    """Vide TOUT le PTA d'une année (suivi + programmes/projets/activités/tâches)."""
    a = db.get_or_404(Annee, ann_id)
    from models import SuiviTache, Programme
    # 1. Suivi d'abord (FK vers taches)
    nb_suivi = SuiviTache.query.filter_by(annee_id=ann_id).delete()
    db.session.flush()
    # 2. Programmes en cascade → projets → activités → tâches
    programmes = requete_pta(ann_id).all()
    nb_prog = len(programmes)
    for prog in programmes:
        db.session.delete(prog)
    db.session.commit()
    log_audit('annee_purge_pta', f"PTA {a.annee} entièrement purgé : {nb_suivi} suivi(s), {nb_prog} programme(s)")
    flash(f'PTA {a.annee} entièrement vidé : {nb_suivi} suivi(s) et '
          f'{nb_prog} programme(s) supprimé(s) (avec projets, activités, tâches). '
          f'L\'année {a.annee} existe toujours et peut être réimportée.', 'success')
    return redirect(url_for('admin.annees'))


@admin_bp.route('/annees/<int:ann_id>/delete', methods=['POST'])
@principal_required
def annee_delete(ann_id):
    a = db.get_or_404(Annee, ann_id)
    if a.actif:
        flash('Impossible de supprimer une année active. Activez une autre année d\'abord.', 'danger')
        return redirect(url_for('admin.annees'))
    if a.programmes:
        flash(f'Impossible : le PTA {a.annee} contient des données. Supprimez d\'abord le contenu.', 'danger')
        return redirect(url_for('admin.annees'))
    db.session.delete(a)
    db.session.commit()
    flash(f'Année {a.annee} supprimée.', 'success')
    return redirect(url_for('admin.annees'))


# ─── Structures externes ────────────────────────────────────────────────────

@admin_bp.route('/structures-externes')
@admin_required
def structures_externes():
    items = StructureExterne.query.order_by(StructureExterne.nom).all()
    return render_template('admin/structures_externes.html', items=items)


@admin_bp.route('/structures-externes/add', methods=['POST'])
@editeur_required
def struct_ext_add():
    nom = request.form.get('nom', '').strip()
    if not nom:
        flash('Le nom est obligatoire.', 'danger')
        return redirect(url_for('admin.structures_externes'))
    se = StructureExterne(
        nom=nom,
        type_org=request.form.get('type_org', '').strip() or None,
        description=request.form.get('description', '').strip() or None,
    )
    db.session.add(se)
    db.session.commit()
    flash(f'Structure externe « {nom} » ajoutée.', 'success')
    return redirect(url_for('admin.structures_externes'))


@admin_bp.route('/structures-externes/<int:se_id>/edit', methods=['POST'])
@editeur_required
def struct_ext_edit(se_id):
    se = db.get_or_404(StructureExterne, se_id)
    nom = request.form.get('nom', '').strip()
    if not nom:
        flash('Le nom est obligatoire.', 'danger')
        return redirect(url_for('admin.structures_externes'))
    se.nom = nom
    se.type_org = request.form.get('type_org', '').strip() or None
    se.description = request.form.get('description', '').strip() or None
    db.session.commit()
    flash(f'Structure externe « {nom} » mise à jour.', 'success')
    return redirect(url_for('admin.structures_externes'))


@admin_bp.route('/structures-externes/<int:se_id>/delete', methods=['POST'])
@editeur_required
def struct_ext_delete(se_id):
    se = db.get_or_404(StructureExterne, se_id)
    nom = se.nom
    db.session.delete(se)
    db.session.commit()
    flash(f'Structure externe « {nom} » supprimée.', 'success')
    return redirect(url_for('admin.structures_externes'))


# ─── Journal d'audit ────────────────────────────────────────────────────────

@admin_bp.route('/journal')
@editeur_required
def journal():
    from models import AuditLog
    import datetime as _dt

    page          = request.args.get('page', 1, type=int)
    filtre_action = request.args.get('action', '').strip()
    filtre_user   = request.args.get('user', '').strip()
    filtre_depuis = request.args.get('depuis', '').strip()

    q = AuditLog.query.order_by(AuditLog.horodatage.desc())

    if filtre_action:
        q = q.filter(AuditLog.action == filtre_action)
    if filtre_user:
        q = q.filter(
            db.or_(
                AuditLog.user_nom.ilike(f'%{filtre_user}%'),
                AuditLog.user_role.ilike(f'%{filtre_user}%'),
            )
        )
    if filtre_depuis:
        try:
            depuis = _dt.datetime.strptime(filtre_depuis, '%Y-%m-%d')
            q = q.filter(AuditLog.horodatage >= depuis)
        except ValueError:
            filtre_depuis = ''

    pagination = q.paginate(page=page, per_page=50, error_out=False)
    entries    = pagination.items

    # Liste de toutes les actions distinctes (pour le filtre)
    actions_dispo = [r[0] for r in db.session.query(AuditLog.action).distinct().order_by(AuditLog.action).all()]

    return render_template('admin/journal.html',
        entries=entries,
        pagination=pagination,
        actions_dispo=actions_dispo,
        filtre_action=filtre_action,
        filtre_user=filtre_user,
        filtre_depuis=filtre_depuis,
    )


@admin_bp.route('/journal/purge', methods=['POST'])
@limiter.limit('5 per minute')
@principal_required
def journal_purge():
    """Purge tout ou partie du journal d'audit."""
    from models import AuditLog
    import datetime as _dt

    mode = request.form.get('mode', 'tout')   # 'tout' ou 'avant'
    nb   = 0

    if mode == 'avant':
        avant_str = request.form.get('avant_date', '').strip()
        try:
            avant = _dt.datetime.strptime(avant_str, '%Y-%m-%d')
            nb = AuditLog.query.filter(AuditLog.horodatage < avant).delete()
        except ValueError:
            flash('Date invalide. Aucune entrée supprimée.', 'danger')
            return redirect(url_for('admin.journal'))
    else:
        nb = AuditLog.query.delete()

    db.session.commit()
    log_audit('journal_purge', f"Journal d'audit purgé : {nb} entrée(s) supprimée(s) (mode : {mode})")
    flash(f'Journal purgé : {nb} entrée(s) supprimée(s).', 'success')
    return redirect(url_for('admin.journal'))


# ─── Sauvegardes PTA ─────────────────────────────────────────────────────────

@admin_bp.route('/backups')
@admin_required
def backups():
    sauvegardes = PTABackup.query.order_by(PTABackup.created_at.desc()).all()
    return render_template('admin/backups.html', sauvegardes=sauvegardes)


@admin_bp.route('/backups/<int:backup_id>/delete', methods=['POST'])
@principal_required
def backup_delete(backup_id):
    b = db.get_or_404(PTABackup, backup_id)
    db.session.delete(b)
    db.session.commit()
    flash('Sauvegarde supprimée définitivement.', 'success')
    return redirect(url_for('admin.backups'))


@admin_bp.route('/backups/<int:backup_id>/restore', methods=['POST'])
@limiter.limit('3 per minute')
@principal_required
def backup_restore(backup_id):
    b = db.get_or_404(PTABackup, backup_id)
    annee = db.session.get(Annee, b.annee_id) if b.annee_id else None
    if not annee:
        flash("L'année PTA de cette sauvegarde n'existe plus. Créez-la d'abord dans Années.", 'danger')
        return redirect(url_for('admin.backups'))

    existing = Programme.query.filter_by(annee_id=annee.id).count()
    if existing > 0:
        flash(
            f"Le PTA {annee.annee} n'est pas vide ({existing} programme(s) existant(s)). "
            f"Réinitialisez d'abord le PTA depuis la page PTA Global avant de restaurer.",
            'danger'
        )
        return redirect(url_for('admin.backups'))

    password = request.form.get('password', '').strip()
    if not current_user.check_password(password):
        flash('Mot de passe incorrect.', 'danger')
        return redirect(url_for('admin.backups'))

    data = json.loads(b.contenu)

    def get_svcs(ids): return [s for s in (db.session.get(Service, i) for i in ids) if s]
    def get_dirs(ids): return [d for d in (db.session.get(Direction, i) for i in ids) if d]
    def get_ses(ids):  return [se for se in (db.session.get(StructureExterne, i) for i in ids) if se]

    nb_acts = 0
    nb_taches = 0

    for pd in data['programmes']:
        prog = Programme(
            annee_id=annee.id, numero=pd['numero'], nom=pd['nom'],
            description=pd.get('description'),
            objectif_specifique=pd.get('objectif_specifique'),
            poids=pd.get('poids', 0),
        )
        db.session.add(prog)
        db.session.flush()

        for pjd in pd['projets']:
            pj = Projet(
                programme_id=prog.id, numero=pjd['numero'], nom=pjd['nom'],
                description=pjd.get('description'), poids=pjd.get('poids', 0),
            )
            db.session.add(pj)
            db.session.flush()

            for ad in pjd['activites']:
                a = Activite(
                    projet_id=pj.id, numero=ad['numero'], nom=ad['nom'],
                    description=ad.get('description'),
                    direction_responsable_id=ad.get('direction_responsable_id'),
                    imputation_budgetaire=ad.get('imputation_budgetaire'),
                    ressources_propres=ad.get('ressources_propres', 0),
                    fadec_affecte=ad.get('fadec_affecte', 0),
                    fadec_non_affecte=ad.get('fadec_non_affecte', 0),
                    autres_partenaires=ad.get('autres_partenaires', 0),
                    autres_fonds=ad.get('autres_fonds', 0),
                    details_financement=ad.get('details_financement'),
                    acteurs_externes=ad.get('acteurs_externes'),
                    periode_debut=ad.get('periode_debut'),
                    periode_fin=ad.get('periode_fin'),
                    mode_execution=ad.get('mode_execution', 'Direct'),
                    poids=ad.get('poids', 0),
                )
                a.services_intervenants = get_svcs(ad.get('services_intervenants', []))
                a.directions_associees  = get_dirs(ad.get('directions_associees', []))
                a.structures_externes   = get_ses(ad.get('structures_externes', []))
                db.session.add(a)
                db.session.flush()
                nb_acts += 1

                for td in ad.get('taches', []):
                    t = Tache(
                        activite_id=a.id, numero=td['numero'],
                        ordre=td.get('ordre', 0), nom=td['nom'],
                        description=td.get('description'),
                        poids=td.get('poids', 0),
                        imputation_budgetaire=td.get('imputation_budgetaire'),
                        ressources_propres=td.get('ressources_propres', 0),
                        fadec_affecte=td.get('fadec_affecte', 0),
                        fadec_non_affecte=td.get('fadec_non_affecte', 0),
                        autres_partenaires=td.get('autres_partenaires', 0),
                        autres_fonds=td.get('autres_fonds', 0),
                        details_financement=td.get('details_financement'),
                        acteurs_externes=td.get('acteurs_externes'),
                        mode_execution=td.get('mode_execution', 'Direct'),
                        service_responsable_id=td.get('service_responsable_id'),
                        direction_responsable_id=td.get('direction_responsable_id'),
                    )
                    t.services_concernes   = get_svcs(td.get('services_concernes', []))
                    t.directions_associees = get_dirs(td.get('directions_associees', []))
                    t.structures_externes  = get_ses(td.get('structures_externes', []))
                    db.session.add(t)
                    nb_taches += 1

    db.session.commit()
    flash(
        f'PTA {annee.annee} restauré depuis la sauvegarde #{b.id} : '
        f'{len(data["programmes"])} programmes, {nb_acts} activités, {nb_taches} tâches.',
        'success'
    )
    return redirect(url_for('pta.global_pta'))


# ─── Emails PTA (rappel saisie + bilan) ──────────────────────────────────────

# Pas de fallback hardcodé : si ADMIN_CC n'est pas configuré, aucune copie supplémentaire.
# Pour activer les copies : ajouter ADMIN_CC=email@domaine.bj dans ~/.pta_backup_config
# ou définir la variable d'environnement PTA_ADMIN_CC sur PythonAnywhere.
_COPIES_FIXES_DEFAUT = []


def _lire_cfg_smtp():
    """Lit la config SMTP depuis les variables d'environnement en priorité,
    puis depuis ~/.pta_backup_config en complément.
    Variables d'env reconnues : PTA_GMAIL_USER, PTA_GMAIL_APP_PASSWORD,
    PTA_ADMIN_CC, PTA_EXPEDITEUR_NOM, PTA_EXPEDITEUR_EMAIL.
    Lève ValueError si GMAIL_USER ou GMAIL_APP_PASSWORD manque partout."""
    import os
    cfg = {}
    # 1. Lire le fichier config (base)
    config_path = os.path.expanduser('~/.pta_backup_config')
    if os.path.exists(config_path):
        with open(config_path, encoding='utf-8') as f:
            for ligne in f:
                ligne = ligne.strip()
                if '=' in ligne and not ligne.startswith('#'):
                    cle, val = ligne.split('=', 1)
                    cfg[cle.strip()] = val.strip()
    # 2. Les variables d'environnement écrasent le fichier (plus sécurisé sur PythonAnywhere)
    _env_map = {
        'PTA_GMAIL_USER':           'GMAIL_USER',
        'PTA_GMAIL_APP_PASSWORD':   'GMAIL_APP_PASSWORD',
        'PTA_ADMIN_CC':             'ADMIN_CC',
        'PTA_EXPEDITEUR_NOM':       'EXPEDITEUR_NOM',
        'PTA_EXPEDITEUR_EMAIL':     'EXPEDITEUR_EMAIL',
    }
    for env_var, cfg_key in _env_map.items():
        val = os.environ.get(env_var, '').strip()
        if val:
            cfg[cfg_key] = val
    for cle in ('GMAIL_USER', 'GMAIL_APP_PASSWORD'):
        if not cfg.get(cle):
            raise ValueError(f"Clé manquante dans config SMTP : {cle}")
    cfg.setdefault('EXPEDITEUR_NOM',   'Jupiter GBOYOU')
    cfg.setdefault('EXPEDITEUR_EMAIL', 'jupiter.gboyou@mairie.bj')
    return cfg


def _get_copies_fixes(cfg):
    """Retourne la liste des adresses de copie fixe (ADMIN_CC).
    Vide si non configuré — pas de fallback hardcodé."""
    admin_cc = cfg.get('ADMIN_CC', '').strip()
    if admin_cc:
        return [a.strip() for a in admin_cc.replace(',', ' ').split() if a.strip()]
    return list(_COPIES_FIXES_DEFAUT)


def _get_destinataires():
    """Emails des utilisateurs actifs ayant une adresse renseignée."""
    return [
        u.email.strip()
        for u in User.query.filter(User.actif.is_(True)).all()
        if u.email and u.email.strip()
    ]


def _emails_par_entite():
    """({direction_id: [emails]}, {service_id: [emails]}) des comptes actifs direction / service."""
    par_dir, par_svc = {}, {}
    for u in User.query.filter(User.actif.is_(True)).all():
        email = (u.email or '').strip()
        if not email:
            continue
        if u.role == 'direction' and u.direction_id:
            par_dir.setdefault(u.direction_id, []).append(email)
        elif u.role == 'service' and u.service_id:
            par_svc.setdefault(u.service_id, []).append(email)
    return par_dir, par_svc


def _ids_formulaire(nom):
    return {int(v) for v in request.form.getlist(nom) if str(v).isdigit()}


_MARQUE     = "Àbójútó"
_LOGO_CID   = "logo_abojuto"
_TRICOLORE  = ("<table width='100%' cellpadding='0' cellspacing='0'><tr>"
               "<td height='4' style='background:#008751;font-size:0;line-height:0;'>&nbsp;</td>"
               "<td height='4' style='background:#FCD116;font-size:0;line-height:0;'>&nbsp;</td>"
               "<td height='4' style='background:#E8112D;font-size:0;line-height:0;'>&nbsp;</td>"
               "</tr></table>")


def _mail_html(cfg, titre, date_str, corps, couleur='#0F3529', largeur=600,
               surtitre="Mairie d'Adja-Ouèrè · PTA &amp; PAI"):
    """Gabarit commun des mails Àbójútó : en-tête logo + marque, corps, pied de page."""
    return f"""<!DOCTYPE html>
<html lang="fr"><head><meta charset="UTF-8"></head>
<body style="margin:0;padding:0;background:#eef2f0;font-family:Arial,Helvetica,sans-serif;">
<table width="100%" cellpadding="0" cellspacing="0" style="background:#eef2f0;padding:28px 0;">
<tr><td align="center">
<table width="{largeur}" cellpadding="0" cellspacing="0"
       style="background:#fff;border-radius:14px;overflow:hidden;box-shadow:0 6px 20px rgba(0,0,0,.08);">
  <tr><td style="background:{couleur};padding:22px 32px 20px;">
    <table cellpadding="0" cellspacing="0"><tr>
      <td style="background:#fff;border-radius:10px;padding:6px 10px;" valign="middle">
        <img src="cid:{_LOGO_CID}" alt="Commune d'Adja-Ouèrè" height="46" style="display:block;height:46px;">
      </td>
      <td style="padding-left:16px;" valign="middle">
        <div style="font-family:Georgia,'Times New Roman',serif;font-size:28px;font-weight:bold;
                    color:#FCD116;line-height:1;">{_MARQUE}</div>
        <div style="margin-top:5px;color:rgba(255,255,255,.75);font-size:11px;
                    letter-spacing:1px;text-transform:uppercase;">{surtitre}</div>
      </td>
    </tr></table>
    <h1 style="margin:20px 0 0;color:#fff;font-size:20px;line-height:1.35;font-weight:bold;">{titre}</h1>
    <p style="margin:4px 0 0;color:rgba(255,255,255,.7);font-size:13px;">{date_str}</p>
  </td></tr>
  <tr><td style="padding:0;">{_TRICOLORE}</td></tr>
  <tr><td style="padding:28px 32px;">{corps}</td></tr>
  <tr><td style="background:#f7f9f8;padding:18px 32px;border-top:1px solid #e3e8e5;">
    <p style="margin:0;color:#8a958f;font-size:11px;line-height:1.7;">
      Message envoyé depuis <strong style="color:#0F3529;">{_MARQUE}</strong>,
      la plateforme de planification et de suivi de la Mairie d'Adja-Ouèrè.<br>
      Émis par : <strong>{cfg['EXPEDITEUR_NOM']}</strong> ·
      <a href="mailto:{cfg['EXPEDITEUR_EMAIL']}" style="color:#006B40;">{cfg['EXPEDITEUR_EMAIL']}</a>
    </p>
    <p style="margin:6px 0 0;color:#8a958f;font-size:11px;">
      &#x1F1E7;&#x1F1EF; République du Bénin &nbsp;·&nbsp; Commune d'Adja-Ouèrè : Union · Travail · Prospérité
    </p>
  </td></tr>
</table>
</td></tr>
</table>
</body></html>"""


def _mail_bouton(url, libelle, couleur='#0F3529'):
    return (f"<div style='text-align:center;margin:24px 0;'>"
            f"<a href='{url}' target='_blank' style='background:{couleur};color:#FCD116;"
            f"text-decoration:none;padding:12px 28px;border-radius:8px;font-weight:700;"
            f"font-size:15px;display:inline-block;'>{libelle}</a></div>")


def _construire_mail(cfg, sujet, texte_brut, html_body, destinataires, copies_fixes, copies=None):
    """Message multipart (texte + HTML) avec le logo de la commune intégré en pièce inline.
    copies=None : envoi groupé (destinataires en Cc). Sinon destinataires en À, copies en Cc.
    Les copies fixes passent uniquement par l'enveloppe SMTP (copie cachée réelle)."""
    import os
    from email.header import Header
    from email.mime.image import MIMEImage
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText
    from email.utils import formataddr
    from flask import current_app

    msg = MIMEMultipart('related')
    msg['From']     = formataddr((f"{_MARQUE} · Mairie d'Adja-Ouèrè", cfg['GMAIL_USER']), charset='utf-8')
    if copies is None:
        msg['To'] = cfg['GMAIL_USER']
        msg['Cc'] = ', '.join(destinataires)
    else:
        msg['To'] = ', '.join(destinataires)
        if copies:
            msg['Cc'] = ', '.join(copies)
    msg['Subject']  = Header(sujet, 'utf-8')
    msg['Reply-To'] = cfg['EXPEDITEUR_EMAIL']

    alt = MIMEMultipart('alternative')
    alt.attach(MIMEText(texte_brut + f"\n\n---\n{_MARQUE} · Mairie d'Adja-Ouèrè", 'plain', 'utf-8'))
    alt.attach(MIMEText(html_body, 'html', 'utf-8'))
    msg.attach(alt)

    logo = os.path.join(current_app.root_path, 'static', 'img', 'logo_commune.png')
    if os.path.exists(logo):
        with open(logo, 'rb') as f:
            img = MIMEImage(f.read(), _subtype='png')
        img.add_header('Content-ID', f'<{_LOGO_CID}>')
        img.add_header('Content-Disposition', 'inline', filename='logo_commune.png')
        msg.attach(img)
    return msg


def _envoyer_smtp(cfg, msg, destinataires, copies_fixes=None, copies=None):
    """Envoie via Gmail SMTP. copies_fixes : liste d'adresses toujours en copie cachée."""
    import smtplib
    if copies_fixes is None:
        copies_fixes = _get_copies_fixes(cfg)
    tous = list(dict.fromkeys([cfg['GMAIL_USER']] + destinataires + (copies or []) + copies_fixes))
    with smtplib.SMTP('smtp.gmail.com', 587) as srv:
        srv.ehlo()
        srv.starttls()
        srv.login(cfg['GMAIL_USER'], cfg['GMAIL_APP_PASSWORD'])
        srv.sendmail(cfg['GMAIL_USER'], tous, msg.as_string())


@admin_bp.route('/rappel-saisie', methods=['POST'])
@limiter.limit('20 per hour')
@editeur_required
def rappel_saisie():
    """Envoie un rappel de saisie aux directions / services choisis, avec copies choisies."""
    import datetime
    from utils import get_annee

    annee       = get_annee()
    annee_label = annee.annee if annee else datetime.date.today().year
    date_str    = datetime.date.today().strftime('%d/%m/%Y')
    plateforme  = request.host_url.rstrip('/')

    par_dir, par_svc = _emails_par_entite()
    dest_dir = _ids_formulaire('dest_dir')
    dest_svc = _ids_formulaire('dest_svc')
    cc_dir   = _ids_formulaire('cc_dir')
    cc_svc   = _ids_formulaire('cc_svc')
    if request.form.get('cc_tutelle'):
        cc_dir |= {s.direction_id for s in Service.query.filter(Service.id.in_(dest_svc)).all()}

    def _emails(ids_dir, ids_svc):
        out = [e for i in sorted(ids_dir) for e in par_dir.get(i, [])]
        out += [e for i in sorted(ids_svc) for e in par_svc.get(i, [])]
        return list(dict.fromkeys(out))

    destinataires = _emails(dest_dir, dest_svc)
    copies = [e for e in _emails(cc_dir, cc_svc) if e not in destinataires]
    if not destinataires:
        flash("Aucune adresse email trouvée pour les destinataires choisis. "
              "Sélectionnez au moins une direction ou un service dont le compte a un email.", 'warning')
        return redirect(url_for('admin.index'))

    try:
        cfg = _lire_cfg_smtp()
    except FileNotFoundError:
        flash("Fichier de configuration email introuvable (~/.pta_backup_config).", 'danger')
        return redirect(url_for('admin.index'))
    except ValueError as e:
        flash(f"Configuration email incomplète : {e}", 'danger')
        return redirect(url_for('admin.index'))
    except OSError as e:
        flash(f"Impossible de lire la configuration email : {e}", 'danger')
        return redirect(url_for('admin.index'))

    copies_fixes = _get_copies_fixes(cfg)
    sujet = f"[{_MARQUE} · PTA {annee_label}] Rappel : renseignez le suivi de votre PTA"

    etapes = [
        f"Connectez-vous à <strong>{_MARQUE}</strong> avec votre identifiant.",
        "Ouvrez le menu (bouton <strong>☰</strong> en haut à gauche), puis "
        "<strong>Suivi &amp; Évaluation PTA</strong>.",
        "Sélectionnez l'onglet du trimestre concerné (<strong>T1</strong> à <strong>T4</strong>), "
        "puis cliquez sur le bouton jaune <strong>Faire le suivi de mon PTA</strong>.",
        "Pour chaque tâche, indiquez le statut (<em>Exécutée</em>, <em>En cours</em> ou "
        "<em>Non exécutée</em>). Si la tâche est en cours, précisez son taux d'avancement. "
        "Ajoutez vos observations ou difficultés si nécessaire.",
        "Cliquez sur <strong>Valider les saisies</strong> (en haut ou en bas du tableau). "
        "Sans cette validation, rien n'est enregistré.",
    ]
    lignes_etapes = ''.join(
        f"<tr><td valign='top' style='padding:0 12px 14px 0;'>"
        f"<div style='width:26px;height:26px;border-radius:13px;background:#0F3529;color:#FCD116;"
        f"font-weight:bold;font-size:13px;line-height:26px;text-align:center;'>{i}</div></td>"
        f"<td valign='top' style='padding:3px 0 14px;color:#374151;font-size:14px;line-height:1.6;'>{e}</td></tr>"
        for i, e in enumerate(etapes, 1)
    )

    corps = f"""
    <p style="margin:0 0 16px;color:#374151;">Madame, Monsieur,</p>
    <p style="margin:0 0 20px;color:#374151;line-height:1.7;">
      Dans le cadre du suivi et de l'évaluation du Plan de Travail Annuel (PTA) {annee_label}
      de la Mairie d'Adja-Ouèrè, nous vous prions de bien vouloir renseigner l'état
      d'avancement de vos tâches sur <strong>{_MARQUE}</strong>.
    </p>
    <p style="margin:0 0 20px;color:#374151;line-height:1.7;">
      Nous comptons sur votre diligence habituelle pour que ce renseignement soit
      effectué dans les meilleurs délais.
    </p>
    <p style="margin:0 0 12px;color:#0F3529;font-weight:bold;font-size:15px;">Comment procéder</p>
    <table cellpadding="0" cellspacing="0" style="margin:0 0 8px;">{lignes_etapes}</table>
    {_mail_bouton(plateforme, f"Ouvrir {_MARQUE}")}
    <p style="margin:16px 0 0;color:#6b7280;font-size:13px;line-height:1.6;">
      Pour toute difficulté de connexion ou de saisie, contactez l'administrateur du système.
    </p>"""
    html_body = _mail_html(cfg, "Rappel : saisie du suivi de votre PTA", date_str, corps,
                           surtitre=f"Mairie d'Adja-Ouèrè · PTA {annee_label}")

    import re as _re
    texte_brut = (
        f"Rappel : saisie du suivi de votre PTA {annee_label}\n"
        f"Date : {date_str}\n\n"
        f"Madame, Monsieur,\n\n"
        f"Dans le cadre du suivi et de l'évaluation du PTA {annee_label}, nous vous prions de bien "
        f"vouloir renseigner l'état d'avancement de vos tâches sur {_MARQUE} ({plateforme}).\n\n"
        "Nous comptons sur votre diligence habituelle pour que ce renseignement soit effectué "
        "dans les meilleurs délais.\n\n"
        "Comment procéder :\n"
        + '\n'.join(f"{i}. {_re.sub(r'<[^>]+>', '', e).replace('&amp;', '&')}"
                    for i, e in enumerate(etapes, 1))
    )

    msg = _construire_mail(cfg, sujet, texte_brut, html_body, destinataires, copies_fixes,
                           copies=copies)

    try:
        _envoyer_smtp(cfg, msg, destinataires, copies_fixes, copies=copies)
    except Exception as e:
        flash(f"Erreur lors de l'envoi : {e}", 'danger')
        return redirect(url_for('admin.index'))

    codes = sorted({d.code for d in Direction.query.filter(Direction.id.in_(dest_dir)).all()}
                   | {s.code for s in Service.query.filter(Service.id.in_(dest_svc)).all()})
    log_audit('rappel_saisie',
              f"Rappel de saisie PTA {annee_label} : {', '.join(codes)} "
              f"({len(destinataires)} destinataire(s), {len(copies)} en copie)")
    flash(f"Rappel envoyé à {len(destinataires)} destinataire(s)"
          f"{f', {len(copies)} en copie' if copies else ''}.", 'success')
    return redirect(url_for('admin.index'))


@admin_bp.route('/purge-suivi', methods=['POST'])
@limiter.limit('5 per minute')
@principal_required
def purge_suivi():
    """Supprime tous les suivis de l'année active (tests uniquement).
    Ramène tout à non_execute et efface toutes les observations."""
    from utils import get_annee
    annee = get_annee()
    if not annee:
        flash("Aucune année PTA active.", 'warning')
        return redirect(url_for('admin.index'))

    nb = SuiviTache.query.filter_by(annee_id=annee.id).count()
    SuiviTache.query.filter_by(annee_id=annee.id).delete()
    db.session.commit()

    log_audit('purge_suivi',
              f"Purge du suivi PTA {annee.annee} : {nb} enregistrement(s) supprimé(s)")
    flash(
        f"Purge effectuée : {nb} suivi(s) supprimé(s) pour l'année {annee.annee}. "
        "Tout est revenu à « Non exécuté ».",
        'success'
    )
    return redirect(url_for('admin.index'))


@admin_bp.route('/purge-pei', methods=['POST'])
@limiter.limit('5 per minute')
@principal_required
def purge_pei():
    """Remet à zéro tous les montants, taux physique et observations du PEI
    pour l'année active (tests uniquement)."""
    from utils import get_annee
    annee = get_annee()
    if not annee:
        flash("Aucune année PTA active.", 'warning')
        return redirect(url_for('admin.index'))

    # Récupère tous les PeiActivite liés à l'année active
    pei_records = (PeiActivite.query
        .join(Activite, PeiActivite.activite_id == Activite.id)
        .join(Projet,   Activite.projet_id == Projet.id)
        .join(Programme, Projet.programme_id == Programme.id)
        .filter(Programme.annee_id == annee.id)
        .all())

    nb = len(pei_records)
    for pei in pei_records:
        pei.montant_engage    = 0.0
        pei.montant_mandate   = 0.0
        pei.montant_paye      = 0.0
        pei.engage_verrouille = False
        pei.taux_physique     = 0.0
        pei.observations      = None
    db.session.commit()

    log_audit('purge_pei',
              f"Purge du PEI {annee.annee} : {nb} activité(s) remise(s) à zéro")
    flash(
        f"Purge PAI effectuée : {nb} activité(s) remise(s) à zéro pour l'année {annee.annee}. "
        "Montants, taux et observations effacés.",
        'success'
    )
    return redirect(url_for('admin.index'))


@admin_bp.route('/bilan-pta', methods=['POST'])
@limiter.limit('5 per hour')
@editeur_required
def bilan_pta():
    """Envoie le bilan global PTA — taux et statuts identiques à l'interface Suivi."""
    import datetime
    from utils import get_annee

    annee       = get_annee()
    annee_label = annee.annee if annee else datetime.date.today().year
    date_str    = datetime.date.today().strftime('%d/%m/%Y')

    destinataires = _get_destinataires()
    if not destinataires:
        flash("Aucun utilisateur actif n'a d'adresse email renseignée.", 'warning')
        return redirect(url_for('admin.index'))

    try:
        cfg = _lire_cfg_smtp()
    except FileNotFoundError:
        flash("Fichier de configuration email introuvable (~/.pta_backup_config).", 'danger')
        return redirect(url_for('admin.index'))
    except ValueError as e:
        flash(f"Configuration email incomplète : {e}", 'danger')
        return redirect(url_for('admin.index'))
    except OSError as e:
        flash(f"Impossible de lire la configuration email : {e}", 'danger')
        return redirect(url_for('admin.index'))

    if not annee:
        flash("Aucune année PTA active.", 'warning')
        return redirect(url_for('admin.index'))

    copies_fixes = _get_copies_fixes(cfg)

    # ── Données — mêmes fonctions que le tableau de bord admin ──────────────────
    from suivi.routes import _compute_pta_global
    from dashboard.routes import (_compute_synthese_admin, _compute_dashboard_stats,
                                   _cibles_global)

    data     = _compute_pta_global(annee)
    cibles   = _cibles_global(annee)
    stats    = _compute_dashboard_stats(annee, data, cibles, with_nature=True)
    synthese = _compute_synthese_admin(annee)   # [{direction, taux, nb_taches, services:[...]}]

    # Global (trim=0)
    taux_global = round(stats['suivi'][0]['taux'], 1)
    glob_a      = stats['suivi'][0]['activites']   # execute / en_cours / non_execute / total
    total_glob  = glob_a['total']

    # Nature (trim=0)
    def _nat_a(key):
        d = stats.get(key) or {}
        return d.get(0, {}).get('activites',
                                {'execute': 0, 'en_cours': 0, 'non_execute': 0, 'total': 0})
    def _nat_t(key):
        d = stats.get(key) or {}
        return round(d.get(0, {}).get('taux', 0.0), 1)

    inv_a    = _nat_a('suivi_invest');  taux_inv = _nat_t('suivi_invest')
    fct_a    = _nat_a('suivi_fonct');   taux_fct = _nat_t('suivi_fonct')

    # ── HTML ──────────────────────────────────────────────────────────────────
    VERT   = '#16a34a'
    ORANGE = '#f59e0b'
    ROUGE  = '#dc2626'

    def _coul(t): return VERT if t >= 75 else (ORANGE if t >= 30 else ROUGE)

    def _cel_taux(t):
        return (f"<td style='padding:8px 10px;text-align:center;border-bottom:1px solid #e5e7eb;"
                f"font-weight:700;font-size:15px;color:{_coul(t)};'>{t}%</td>")

    def _lig_nat(acts, taux):
        tot = acts['total']
        return (
            f"<td style='padding:8px 10px;text-align:center;border-bottom:1px solid #e5e7eb;'>"
            f"<b style='color:{VERT};'>{acts['execute']}</b></td>"
            f"<td style='padding:8px 10px;text-align:center;border-bottom:1px solid #e5e7eb;'>"
            f"<b style='color:{ORANGE};'>{acts['en_cours']}</b></td>"
            f"<td style='padding:8px 10px;text-align:center;border-bottom:1px solid #e5e7eb;'>"
            f"<b style='color:{ROUGE};'>{acts['non_execute']}</b></td>"
            f"<td style='padding:8px 10px;text-align:center;border-bottom:1px solid #e5e7eb;"
            f"font-weight:700;'>{tot}</td>"
            + _cel_taux(taux)
        )

    TH  = "style='padding:10px;text-align:left;'"
    THC = "style='padding:10px;text-align:center;'"
    TS  = ("width='100%' cellpadding='0' cellspacing='0' "
           "style='border-collapse:collapse;font-size:13px;margin:10px 0 18px;'")
    VIDE3 = "<tr><td colspan='3' style='padding:10px;color:#9ca3af;text-align:center;'>—</td></tr>"
    VIDE5 = "<tr><td colspan='5' style='padding:10px;color:#9ca3af;text-align:center;'>—</td></tr>"

    # Table directions — Code | Direction | Taux
    # Seules les directions ayant au moins une tâche dans le PTA (nb_taches > 0)
    lig_d = ''.join(
        f"<tr>"
        f"<td style='padding:8px 10px;border-bottom:1px solid #e5e7eb;"
        f"font-weight:700;color:#1e3a5f;'>{item['direction'].code}</td>"
        f"<td style='padding:8px 10px;border-bottom:1px solid #e5e7eb;'>"
        f"{item['direction'].nom}</td>"
        + _cel_taux(round(item['taux'], 1)) + "</tr>"
        for item in synthese if item['nb_taches'] > 0
    ) or VIDE3
    entete_d = (f"<tr style='background:#1e3a5f;color:#fff;'>"
                f"<th {TH}>Code</th><th {TH}>Direction</th>"
                f"<th {THC}>Taux d'exécution</th></tr>")
    bloc_dir = f"<table {TS}>{entete_d}<tbody>{lig_d}</tbody></table>"

    # Table services — Code | Service (Dir) | Taux
    # Seuls les services ayant au moins une tâche dans le PTA (nb_taches > 0)
    lig_s = ''.join(
        f"<tr>"
        f"<td style='padding:8px 10px;border-bottom:1px solid #e5e7eb;"
        f"font-weight:700;color:#0f6f3a;'>{sv['service'].code}</td>"
        f"<td style='padding:8px 10px;border-bottom:1px solid #e5e7eb;'>"
        f"{sv['service'].nom} "
        f"<span style='color:#9ca3af;font-size:11px;'>({item['direction'].code})</span></td>"
        + _cel_taux(round(sv['taux'], 1)) + "</tr>"
        for item in synthese
        for sv in item['services'] if sv['nb_taches'] > 0
    ) or VIDE3
    entete_s = (f"<tr style='background:#0f6f3a;color:#fff;'>"
                f"<th {TH}>Code</th><th {TH}>Service (Direction)</th>"
                f"<th {THC}>Taux d'exécution</th></tr>")
    bloc_svc = f"<table {TS}>{entete_s}<tbody>{lig_s}</tbody></table>"

    # Table nature — ✅ En cours ⏸ Total Taux
    entete_nat = (
        f"<tr style='background:#1e3a5f;color:#fff;'>"
        f"<th {TH}>Nature</th>"
        f"<th {THC}>✅ Exéc.</th><th {THC}>🔄 En cours</th>"
        f"<th {THC}>⏸ Non exéc.</th><th {THC}>Total</th>"
        f"<th {THC}>Taux</th></tr>"
    )
    lig_n = (
        f"<tr><td style='padding:8px 10px;border-bottom:1px solid #e5e7eb;"
        f"font-weight:700;'>Investissement</td>" + _lig_nat(inv_a, taux_inv) + "</tr>"
        f"<tr><td style='padding:8px 10px;border-bottom:1px solid #e5e7eb;"
        f"font-weight:700;'>Fonctionnement</td>" + _lig_nat(fct_a, taux_fct) + "</tr>"
    )
    bloc_nat = f"<table {TS}>{entete_nat}<tbody>{lig_n}</tbody></table>"

    # Synthèse globale
    pct_e = round(glob_a['execute']     / total_glob * 100, 1) if total_glob else 0
    pct_c = round(glob_a['en_cours']    / total_glob * 100, 1) if total_glob else 0
    pct_n = round(glob_a['non_execute'] / total_glob * 100, 1) if total_glob else 0
    bloc_glob = (
        f"<table width='100%' cellpadding='0' cellspacing='0' "
        f"style='border-collapse:collapse;font-size:14px;'>"
        f"<tr style='background:#f1f5f9;'>"
        f"<td style='padding:12px;font-weight:700;color:#1e3a5f;'>"
        f"TOTAL : {total_glob} activité(s)</td>"
        f"<td style='padding:12px;text-align:center;'>"
        f"<span style='color:{VERT};font-weight:700;font-size:18px;'>{glob_a['execute']}</span><br>"
        f"<span style='font-size:11px;color:#6b7280;'>exécutée(s) ({pct_e}%)</span></td>"
        f"<td style='padding:12px;text-align:center;'>"
        f"<span style='color:{ORANGE};font-weight:700;font-size:18px;'>{glob_a['en_cours']}</span><br>"
        f"<span style='font-size:11px;color:#6b7280;'>en cours ({pct_c}%)</span></td>"
        f"<td style='padding:12px;text-align:center;'>"
        f"<span style='color:{ROUGE};font-weight:700;font-size:18px;'>{glob_a['non_execute']}</span><br>"
        f"<span style='font-size:11px;color:#6b7280;'>non exéc. ({pct_n}%)</span></td>"
        f"<td style='padding:12px;text-align:center;'>"
        f"<span style='color:{_coul(taux_global)};font-weight:700;font-size:22px;'>"
        f"{taux_global}%</span><br>"
        f"<span style='font-size:11px;color:#6b7280;'>taux global</span></td>"
        f"</tr></table>"
    )

    def sec(t): return f"<h2 style='margin:20px 0 6px;color:#1e3a5f;font-size:15px;border-bottom:2px solid #1e3a5f;padding-bottom:5px;'>{t}</h2>"

    corps = f"""
    <p style="margin:0 0 20px;color:#374151;line-height:1.6;">
      Bonjour,<br>Bilan d'avancement du PTA {annee_label} à la date du <strong>{date_str}</strong>.
    </p>
    {sec('📊 Activités par direction')}
    {bloc_dir}
    {sec('🏢 Activités par service')}
    {bloc_svc}
    {sec('🏷️ Activités par nature')}
    {bloc_nat}
    {sec('🌐 Synthèse globale')}
    {bloc_glob}
    <p style="margin:20px 0 0;color:#6b7280;font-size:12px;">
      Pour le détail complet, connectez-vous à {_MARQUE}.
    </p>"""
    html_body = _mail_html(cfg, "Bilan global du PTA", f"État au {date_str}", corps,
                           largeur=720, surtitre=f"Mairie d'Adja-Ouèrè · PTA {annee_label}")

    texte_brut = (
        f"Bilan global du PTA {annee_label}, état au {date_str}\n\n"
        f"GLOBAL : {total_glob} activité(s) | {glob_a['execute']} exéc. | "
        f"{glob_a['en_cours']} en cours | {glob_a['non_execute']} non exéc. | Taux : {taux_global}%\n\n"
        "PAR DIRECTION :\n"
        + '\n'.join(
            f"  {item['direction'].code} · {item['direction'].nom} : taux {round(item['taux'], 1)}%"
            for item in synthese if item['nb_taches'] > 0
        )
        + "\n\nPAR SERVICE :\n"
        + '\n'.join(
            f"  {sv['service'].code} ({item['direction'].code}) · "
            f"{sv['service'].nom} : Taux {round(sv['taux'], 1)}%"
            for item in synthese
            for sv in item['services'] if sv['nb_taches'] > 0
        )
        + f"\n\nPAR NATURE :\n"
        f"  Investissement : {inv_a['total']} act. | Taux {taux_inv}%\n"
        f"  Fonctionnement : {fct_a['total']} act. | Taux {taux_fct}%"
    )

    sujet = f"[{_MARQUE} · PTA {annee_label}] Bilan global au {date_str}"

    msg = _construire_mail(cfg, sujet, texte_brut, html_body, destinataires, copies_fixes)

    try:
        _envoyer_smtp(cfg, msg, destinataires, copies_fixes)
    except Exception as e:
        flash(f"Erreur lors de l'envoi : {e}", 'danger')
        return redirect(url_for('admin.index'))

    log_audit('bilan_pta',
              f"Bilan PTA {annee_label} envoyé à {len(destinataires)} destinataire(s) "
              f"— {total_glob} activité(s), taux global {taux_global}%")
    flash(f"Bilan envoyé à {len(destinataires)} destinataire(s) + {len(copies_fixes)} copie(s).", 'success')
    return redirect(url_for('admin.index'))


@admin_bp.route('/notifier-pta-pai', methods=['POST'])
@limiter.limit('10 per hour')
@editeur_required
def notifier_pta_pai():
    """Informe tous les utilisateurs d'une mise à jour PTA/PAI ou de la disponibilité d'une nouvelle année."""
    import datetime
    from models import Annee

    type_notif  = request.form.get('type_notif', 'maj')   # 'maj' ou 'nouvelle_annee'
    annee_id    = request.form.get('annee_id', type=int)
    annee       = db.session.get(Annee, annee_id) if annee_id else None
    if not annee:
        flash("Année introuvable.", 'danger')
        return redirect(url_for('admin.index'))

    annee_label = annee.annee
    date_str    = datetime.date.today().strftime('%d/%m/%Y')
    plateforme  = request.host_url.rstrip('/')

    destinataires = _get_destinataires()
    if not destinataires:
        flash("Aucun utilisateur actif n'a d'adresse email renseignée.", 'warning')
        return redirect(url_for('admin.index'))

    try:
        cfg = _lire_cfg_smtp()
    except FileNotFoundError:
        flash("Fichier de configuration email introuvable (~/.pta_backup_config).", 'danger')
        return redirect(url_for('admin.index'))
    except (ValueError, OSError) as e:
        flash(f"Configuration email : {e}", 'danger')
        return redirect(url_for('admin.index'))

    copies_fixes = _get_copies_fixes(cfg)

    if type_notif == 'nouvelle_annee':
        sujet      = f"[{_MARQUE} · PTA {annee_label}] Nouveau PTA et PAI {annee_label} disponibles"
        titre_mail = f"Nouveau PTA & PAI {annee_label} disponibles"
        intro      = (
            f"Le Plan de Travail Annuel (PTA) et le Plan Annuel d'Investissement (PAI) "
            f"pour l'exercice <strong>{annee_label}</strong> sont désormais disponibles sur la plateforme."
        )
        badge_couleur = '#0f6f3a'
        badge_texte   = f'Nouvelle année {annee_label}'
    else:
        sujet      = f"[{_MARQUE} · PTA {annee_label}] PTA et PAI {annee_label} mis à jour"
        titre_mail = f"PTA & PAI {annee_label} mis à jour"
        intro      = (
            f"Le Plan de Travail Annuel (PTA) et le Plan Annuel d'Investissement (PAI) "
            f"de l'exercice <strong>{annee_label}</strong> ont été mis à jour. "
            "Connectez-vous pour consulter les dernières informations."
        )
        badge_couleur = '#1e3a5f'
        badge_texte   = f'Mise à jour {annee_label}'

    corps = f"""
    <p style="margin:0 0 16px;color:#374151;">Madame, Monsieur,</p>
    <p style="margin:0 0 20px;color:#374151;line-height:1.7;">{intro}</p>
    {_mail_bouton(plateforme, f"Se connecter à {_MARQUE}", badge_couleur)}
    <p style="margin:16px 0 0;color:#6b7280;font-size:12px;">
      En cas de difficulté de connexion, contactez votre administrateur.
    </p>"""
    html_body = _mail_html(cfg, titre_mail, date_str, corps, couleur=badge_couleur)

    texte_brut = (
        f"{titre_mail}\n\n"
        f"{intro.replace('<strong>','').replace('</strong>','')}\n\n"
        f"Connectez-vous : {plateforme}"
    )

    msg = _construire_mail(cfg, sujet, texte_brut, html_body, destinataires, copies_fixes)

    try:
        _envoyer_smtp(cfg, msg, destinataires, copies_fixes)
    except Exception as e:
        flash(f"Erreur lors de l'envoi : {e}", 'danger')
        return redirect(url_for('admin.index'))

    log_audit('notifier_pta_pai',
              f"Notification PTA/PAI «{type_notif}» — année {annee_label} — "
              f"{len(destinataires)} destinataire(s)")
    flash(f"Notification envoyée à {len(destinataires)} utilisateur(s) + {len(copies_fixes)} copie(s).", 'success')
    return redirect(url_for('admin.index'))


# ─── Mode maintenance ────────────────────────────────────────────────────────

@admin_bp.route('/maintenance/activer', methods=['POST'])
@principal_required
def maintenance_activer():
    import maintenance
    maintenance.activer(request.form.get('message', ''), request.form.get('retour', ''),
                        par=f"{current_user.prenom} {current_user.nom}")
    log_audit('maintenance_on', "Mode maintenance activé")
    flash("Mode maintenance activé : les autres utilisateurs n'ont plus accès à la plateforme.", 'warning')
    return redirect(url_for('admin.index'))


@admin_bp.route('/maintenance/desactiver', methods=['POST'])
@principal_required
def maintenance_desactiver():
    import maintenance
    maintenance.desactiver()
    log_audit('maintenance_off', "Mode maintenance désactivé")
    flash("Mode maintenance désactivé : la plateforme est de nouveau ouverte à tous.", 'success')
    return redirect(request.referrer if (request.referrer or '').startswith(request.host_url)
                    else url_for('admin.index'))


# ─── Archives : copies figées du PTA et du suivi ─────────────────────────────

def _dossier_archives():
    import os
    from flask import current_app
    d = os.path.join(current_app.instance_path, 'archives')
    os.makedirs(d, exist_ok=True)
    return d


@admin_bp.route('/archives')
@editeur_required
def archives():
    import os
    from models import Archive
    liste = Archive.query.order_by(Archive.annee_label.desc(), Archive.created_at.desc()).all()
    dossier = _dossier_archives()
    from utils import get_annee
    presents = {a.id for a in liste if os.path.exists(os.path.join(dossier, a.fichier))}
    annee = get_annee()
    existantes = {f"{a.type_archive}|{a.trimestre if a.trimestre is not None else ''}":
                  a.created_at.strftime('%d/%m/%Y à %Hh%M')
                  for a in liste if annee and a.annee_label == annee.annee}
    par_annee = {}
    for a in liste:
        par_annee.setdefault(a.annee_label, []).append(a)
    return render_template('admin/archives.html', par_annee=par_annee, presents=presents,
                           types=Archive.TYPES, avec_trimestre=Archive.AVEC_TRIMESTRE,
                           annee_courante=annee, existantes=existantes)


@admin_bp.route('/archives/creer', methods=['POST'])
@limiter.limit('10 per hour')
@editeur_required
def archive_creer():
    import os, re, io
    from datetime import datetime, timezone
    from models import Archive
    from utils import get_annee
    from exportation.routes import _build_workbook, _build_suivi_workbook

    annee = get_annee()
    if not annee:
        flash("Aucune année PTA sélectionnée.", 'danger')
        return redirect(url_for('admin.archives'))

    type_archive = request.form.get('type_archive', '')
    if type_archive not in Archive.TYPES:
        flash("Type d'archive inconnu.", 'danger')
        return redirect(url_for('admin.archives'))
    # Pour le suivi, le point d'exécution et le budget, le trimestre n'est qu'un repère :
    # on archive toujours l'état complet, tel qu'il est à la fin de ce trimestre.
    trimestre = None
    if type_archive in Archive.AVEC_TRIMESTRE:
        trimestre = request.form.get('trimestre', 0, type=int)
        if trimestre not in (1, 2, 3, 4):
            flash("Choisissez le trimestre de référence (T1 à T4).", 'warning')
            return redirect(url_for('admin.archives'))

    directions = Direction.query.order_by(Direction.nom).all()
    services   = Service.query.order_by(Service.nom).all()
    try:
        if type_archive in ('pta_initial', 'pta_revise'):
            wb = _build_workbook(annee, include_global=True, include_recap=True,
                                 directions=directions, services=services)
            contenu, defaut = _octets(wb), f"{Archive.TYPES[type_archive]} {annee.annee}"
        elif type_archive == 'suivi':
            # Vue globale : toutes les tâches, avec leur dernier état saisi (aucun filtre)
            wb = _build_suivi_workbook(annee, trimestre=0, include_global=True,
                                       dirs=directions, svcs=services)
            contenu = _octets(wb)
            defaut = f"Suivi & Évaluation du PTA {annee.annee} (situation au T{trimestre})"
        elif type_archive in ('pai_initial', 'pai_revise'):
            from pai.routes import export_excel as _export_pai
            contenu = _capturer_export(_export_pai, '/pai/export')
            defaut = f"{Archive.TYPES[type_archive]} {annee.annee}"
        elif type_archive == 'pei':
            from pei.routes import export_excel as _export_pei
            contenu = _capturer_export(_export_pei, '/pei/export')
            defaut = f"Point d'exécution du PAI {annee.annee} (situation au T{trimestre})"
        else:  # budget
            from budget.routes import export_excel as _export_budget
            # Les quatre trimestres, tels qu'ils sont renseignés à cette date
            contenu = _capturer_export(_export_budget, '/budget/export', {'t': [1, 2, 3, 4]})
            defaut = f"Exécution du budget {annee.annee} (situation au T{trimestre})"
    except Exception as e:
        flash(f"La création de l'archive a échoué : {e}", 'danger')
        return redirect(url_for('admin.archives'))

    from openpyxl import load_workbook
    nb_feuilles = len(load_workbook(io.BytesIO(contenu), read_only=True).sheetnames)

    libelle = (request.form.get('libelle', '').strip() or defaut)[:200]
    horodatage = datetime.now().strftime('%Y-%m-%d_%Hh%M%S')
    base = re.sub(r'[^A-Za-z0-9]+', '_',
                  f"{type_archive}_{annee.annee}" + (f"_T{trimestre}" if trimestre else '')).strip('_')
    fichier = f"{annee.annee}/{base}_{horodatage}.xlsx"
    chemin = os.path.join(_dossier_archives(), fichier)
    os.makedirs(os.path.dirname(chemin), exist_ok=True)
    with open(chemin, 'wb') as f:
        f.write(contenu)

    # Une seule archive par (année, type, trimestre) : un nouvel archivage remplace l'ancien
    a = Archive.query.filter_by(annee_label=annee.annee, type_archive=type_archive,
                                trimestre=trimestre).first()
    remplace = a is not None
    if remplace:
        if a.fichier != fichier:
            try:
                os.remove(os.path.join(_dossier_archives(), a.fichier))
            except OSError:
                pass
    else:
        a = Archive(annee_label=annee.annee, type_archive=type_archive, trimestre=trimestre)
        db.session.add(a)
    a.libelle, a.fichier, a.taille, a.nb_feuilles = libelle, fichier, len(contenu), nb_feuilles
    a.created_at, a.created_by_id = datetime.now(), current_user.id   # heure locale (TZ du serveur)
    db.session.commit()
    log_audit('archive_remplacee' if remplace else 'archive_creee',
              f"Archive {'remplacée' if remplace else 'créée'} : {libelle} ({nb_feuilles} feuille(s))")
    flash(f"Archive « {libelle} » {'remplacée par la version actuelle' if remplace else 'créée'}. "
          "Elle restera téléchargeable telle quelle, même si les données évoluent ensuite.", 'success')
    return redirect(url_for('admin.archives'))


def _octets(wb):
    import io
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def _capturer_export(vue, chemin, query=None):
    """Appelle une vue d'export Excel existante et récupère le fichier produit,
    avec l'utilisateur et l'année de la requête en cours."""
    from flask import current_app, session as _s
    from flask_login import login_user
    utilisateur = current_user._get_current_object()
    annee_id, annee_val = _s.get('annee_id'), _s.get('annee')
    with current_app.test_request_context(chemin, query_string=query or {}):
        login_user(utilisateur)
        _s['annee_id'], _s['annee'] = annee_id, annee_val
        resp = vue()
        if resp.status_code != 200:
            raise RuntimeError(f"export indisponible (code {resp.status_code})")
        resp.direct_passthrough = False
        return resp.get_data()


@admin_bp.route('/archives/<int:archive_id>/telecharger')
@editeur_required
def archive_telecharger(archive_id):
    import os, re
    from flask import send_file
    from models import Archive
    a = db.get_or_404(Archive, archive_id)
    chemin = os.path.join(_dossier_archives(), a.fichier)
    if not os.path.exists(chemin):
        flash("Le fichier de cette archive est introuvable sur ce serveur.", 'danger')
        return redirect(url_for('admin.archives'))
    nom = re.sub(r'[^\w\-]+', '_', a.libelle, flags=re.UNICODE).strip('_') + \
        f"_{a.created_at.strftime('%Y-%m-%d')}.xlsx"
    return send_file(chemin, as_attachment=True, download_name=nom,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@admin_bp.route('/archives/<int:archive_id>/supprimer', methods=['POST'])
@principal_required
def archive_supprimer(archive_id):
    import os
    from models import Archive
    a = db.get_or_404(Archive, archive_id)
    try:
        os.remove(os.path.join(_dossier_archives(), a.fichier))
    except OSError:
        pass
    libelle = a.libelle
    db.session.delete(a)
    db.session.commit()
    log_audit('archive_supprimee', f"Archive supprimée : {libelle}")
    flash(f"Archive « {libelle} » supprimée.", 'warning')
    return redirect(url_for('admin.archives'))
