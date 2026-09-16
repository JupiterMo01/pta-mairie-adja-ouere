from flask import render_template, abort, request, redirect, url_for, flash
from flask_login import login_required, current_user
from models import (db, Programme, PaiActivite, PaiProgramme, PaiProjet,
                    PeiActivite, SuiviTache)
from utils import get_annee, log_audit
from . import pei_bp


# ─── Helpers ─────────────────────────────────────────────────────────────────

def _peut_editer():
    if current_user.role == 'admin_editeur':
        return True
    if current_user.role == 'direction' and current_user.direction \
            and current_user.direction.code == 'DAAF':
        return True
    if current_user.role == 'service' and current_user.service \
            and current_user.service.code == 'SBFC':
        return True
    return False


def _taux_sv(sv):
    if sv is None:             return 0.0
    if sv.statut == 'execute': return 100.0
    if sv.statut == 'en_cours': return float(sv.taux_execution or 0)
    return 0.0


def _load_suivi_global(annee_id):
    """Taux le plus récent par tâche (trimestre le plus élevé)."""
    suivis = (SuiviTache.query
              .filter_by(annee_id=annee_id)
              .order_by(SuiviTache.trimestre.desc())
              .all())
    result = {}
    for s in suivis:
        ex = result.get(s.tache_id)
        if ex is None:
            result[s.tache_id] = s
        elif (s.trimestre == ex.trimestre
              and s.service_id is not None
              and ex.service_id is None):
            result[s.tache_id] = s
    return result


def _wavg(pairs):
    """Moyenne pondérée [(valeur, poids)]. Renvoie 0.0 si poids total nul."""
    p = sum(w for _, w in pairs)
    return round(sum(v * w for v, w in pairs) / p, 1) if p else 0.0


def _build_pei_data(annee):
    suivi_map = _load_suivi_global(annee.id)

    # Toutes observations de l'année (tous trimestres)
    all_svs = SuiviTache.query.filter(
        SuiviTache.annee_id == annee.id,
        SuiviTache.observation.isnot(None)
    ).all()
    obs_map = {}
    for sv in all_svs:
        txt = (sv.observation or '').strip()
        if txt:
            obs_map.setdefault(sv.tache_id, []).append(txt)

    programmes = (Programme.query
                  .filter_by(annee_id=annee.id)
                  .order_by(Programme.numero)
                  .all())
    data = []
    gl_budget = gl_engage = gl_mandate = gl_paye = 0.0
    prog_num = 0

    for pg in programmes:
        pg_projets = []
        proj_num = 0

        for pj in pg.projets:
            inv_acts = [a for a in pj.activites
                        if a.type_activite
                        and 'investissement' in a.type_activite.lower()
                        and a.inclure_dans_pai is not False]
            if not inv_acts:
                continue

            proj_num += 1
            act_num = 0
            act_datas = []

            for act in inv_acts:
                act_num += 1

                # Taux physique : moyenne pondérée des tâches (poids PTA)
                task_pairs = []
                for t in act.taches:
                    task_pairs.append((_taux_sv(suivi_map.get(t.id)), t.poids or 0))
                taux_phys = _wavg(task_pairs) if task_pairs else 0.0

                # Observations compilées
                obs_parts = []
                for t in act.taches:
                    obs_parts.extend(obs_map.get(t.id, []))
                obs_compiled = ' ; '.join(obs_parts)

                # PeiActivite (crée si absent)
                pei = PeiActivite.query.filter_by(activite_id=act.id).first()
                if not pei:
                    pei = PeiActivite(activite_id=act.id)
                    db.session.add(pei)

                budget   = act.budget_total
                engage   = pei.montant_engage  or 0.0
                mandate  = pei.montant_mandate or 0.0
                paye     = pei.montant_paye    or 0.0

                taux_eng  = round(engage  / budget  * 100, 1) if budget  else 0.0
                taux_mand = round(mandate / engage  * 100, 1) if engage  else 0.0
                taux_pay  = round(paye    / engage  * 100, 1) if engage  else 0.0

                pai_extra = act.pai_extra
                poids = int(pai_extra.poids_pai or 0) if pai_extra else 0

                gl_budget  += budget
                gl_engage  += engage
                gl_mandate += mandate
                gl_paye    += paye

                act_datas.append({
                    'activite':  act,
                    'act_num':   act_num,
                    'poids':     poids,
                    'taux_phys': taux_phys,
                    'budget':    budget,
                    'engage':    engage,
                    'mandate':   mandate,
                    'paye':      paye,
                    'taux_eng':  taux_eng,
                    'taux_mand': taux_mand,
                    'taux_pay':  taux_pay,
                    'obs':       obs_compiled,
                    'pei':       pei,
                })

            # Projet — agrégation par poids PAI activités
            pj_extra = PaiProjet.query.filter_by(projet_id=pj.id).first()
            pj_poids = int(pj_extra.poids_pai or 0) if pj_extra else 0

            pj_budget  = sum(ad['budget']  for ad in act_datas)
            pj_engage  = sum(ad['engage']  for ad in act_datas)
            pj_mandate = sum(ad['mandate'] for ad in act_datas)
            pj_paye    = sum(ad['paye']    for ad in act_datas)

            pj_taux_phys = _wavg([(ad['taux_phys'], ad['poids']) for ad in act_datas])
            pj_taux_eng  = _wavg([(ad['taux_eng'],  ad['poids']) for ad in act_datas])
            pj_taux_mand = _wavg([(ad['taux_mand'], ad['poids']) for ad in act_datas])
            pj_taux_pay  = _wavg([(ad['taux_pay'],  ad['poids']) for ad in act_datas])

            pg_projets.append({
                'projet':     pj,
                'proj_num':   proj_num,
                'poids':      pj_poids,
                'activites':  act_datas,
                'budget':     pj_budget,
                'engage':     pj_engage,
                'mandate':    pj_mandate,
                'paye':       pj_paye,
                'taux_phys':  pj_taux_phys,
                'taux_eng':   pj_taux_eng,
                'taux_mand':  pj_taux_mand,
                'taux_pay':   pj_taux_pay,
            })

        if not pg_projets:
            continue

        prog_num += 1
        pg_extra = PaiProgramme.query.filter_by(programme_id=pg.id).first()
        pg_poids = int(pg_extra.poids_pai or 0) if pg_extra else 0

        pg_budget  = sum(pjd['budget']  for pjd in pg_projets)
        pg_engage  = sum(pjd['engage']  for pjd in pg_projets)
        pg_mandate = sum(pjd['mandate'] for pjd in pg_projets)
        pg_paye    = sum(pjd['paye']    for pjd in pg_projets)

        pg_taux_phys = _wavg([(pjd['taux_phys'], pjd['poids']) for pjd in pg_projets])
        pg_taux_eng  = _wavg([(pjd['taux_eng'],  pjd['poids']) for pjd in pg_projets])
        pg_taux_mand = _wavg([(pjd['taux_mand'], pjd['poids']) for pjd in pg_projets])
        pg_taux_pay  = _wavg([(pjd['taux_pay'],  pjd['poids']) for pjd in pg_projets])

        data.append({
            'programme':  pg,
            'prog_num':   prog_num,
            'poids':      pg_poids,
            'projets':    pg_projets,
            'budget':     pg_budget,
            'engage':     pg_engage,
            'mandate':    pg_mandate,
            'paye':       pg_paye,
            'taux_phys':  pg_taux_phys,
            'taux_eng':   pg_taux_eng,
            'taux_mand':  pg_taux_mand,
            'taux_pay':   pg_taux_pay,
        })

    db.session.commit()

    gl_taux_phys = _wavg([(pd['taux_phys'], pd['poids']) for pd in data]) if data else 0.0
    gl_taux_eng  = _wavg([(pd['taux_eng'],  pd['poids']) for pd in data]) if data else 0.0
    gl_taux_mand = _wavg([(pd['taux_mand'], pd['poids']) for pd in data]) if data else 0.0
    gl_taux_pay  = _wavg([(pd['taux_pay'],  pd['poids']) for pd in data]) if data else 0.0

    totaux = {
        'budget':     gl_budget,
        'engage':     gl_engage,
        'mandate':    gl_mandate,
        'paye':       gl_paye,
        'taux_phys':  gl_taux_phys,
        'taux_eng':   gl_taux_eng,
        'taux_mand':  gl_taux_mand,
        'taux_pay':   gl_taux_pay,
    }
    return data, totaux


# ─── Routes ──────────────────────────────────────────────────────────────────

@pei_bp.route('/')
@login_required
def index():
    annee = get_annee()
    if not annee:
        abort(404)
    if not annee.actif and current_user.role not in ('admin_editeur', 'admin_lecteur'):
        abort(403)

    data, totaux = _build_pei_data(annee)
    go = request.args.get('go', '')

    return render_template(
        'pei/index.html',
        annee=annee,
        data=data,
        totaux=totaux,
        peut_editer=_peut_editer(),
        go=go,
    )


@pei_bp.route('/print')
@login_required
def print_view():
    annee = get_annee()
    if not annee:
        abort(404)
    if not annee.actif and current_user.role not in ('admin_editeur', 'admin_lecteur'):
        abort(403)
    data, totaux = _build_pei_data(annee)
    return render_template('pei/print.html', annee=annee, data=data, totaux=totaux)


@pei_bp.route('/export')
@login_required
def export_excel():
    annee = get_annee()
    if not annee:
        abort(404)
    if not annee.actif and current_user.role not in ('admin_editeur', 'admin_lecteur'):
        abort(403)

    data, totaux = _build_pei_data(annee)

    import io, openpyxl, os as _os
    from flask import current_app, send_file
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f'PEI {annee.annee}'

    def fill(hex_color):
        return PatternFill('solid', fgColor=hex_color)

    thin  = Side(style='thin')
    bord  = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left   = Alignment(horizontal='left',   vertical='center', wrap_text=True)

    # ── En-tête institutionnel (lignes 1-3) ──────────────────────────────────
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 8
    ws.row_dimensions[3].height = 22
    ws.merge_cells('A1:D3')
    ws.merge_cells('E1:I3')
    _cc = ws['E1']
    _cc.value = "BP 02 Adja-Ouèrè\nTél : +229 01 61 91 96 12\nEmail : contact.adjaouere@mairie.bj"
    _cc.font  = Font(bold=True, size=9)
    _cc.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.merge_cells('J1:P3')
    _static_img = _os.path.join(current_app.root_path, 'static', 'img')
    try:
        from openpyxl.drawing.image import Image as _XLImg
        from PIL import Image as _PILImg
        import io as _imgio
        _H = 52
        _band = _os.path.join(_static_img, 'bandeau.png')
        _logo = _os.path.join(_static_img, 'logo_commune.png')
        if _os.path.exists(_band):
            _buf = _imgio.BytesIO()
            with _PILImg.open(_band) as _p:
                _ow, _oh = _p.size; _p.save(_buf, format='PNG')
            _buf.seek(0); _i = _XLImg(_buf)
            _i.height = _H; _i.width = int(_ow * _H / _oh)
            ws.add_image(_i, 'A1')
        if _os.path.exists(_logo):
            _buf2 = _imgio.BytesIO()
            with _PILImg.open(_logo) as _p2:
                _ow2, _oh2 = _p2.size; _p2.save(_buf2, format='PNG')
            _buf2.seek(0); _i2 = _XLImg(_buf2)
            _lw = int(_ow2 * _H / _oh2); _i2.height = _H; _i2.width = _lw
            try:
                from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
                from openpyxl.drawing.xdr import XDRPositiveSize2D
                _anch = OneCellAnchor()
                _anch._from = AnchorMarker(col=15, colOff=-int(_lw*9525), row=0, rowOff=0)
                _anch.ext  = XDRPositiveSize2D(int(_lw*9525), int(_H*9525))
                _i2.anchor = _anch; ws.add_image(_i2)
            except Exception:
                ws.add_image(_i2, 'O1')
    except Exception:
        pass

    # ── Titre (ligne 4) ───────────────────────────────────────────────────────
    ws.merge_cells('A4:P4')
    ws['A4'].value = f"POINT D'EXÉCUTION DU PLAN ANNUEL D'INVESTISSEMENT (PAI) — Exercice {annee.annee}"
    ws['A4'].font  = Font(bold=True, size=13)
    ws['A4'].alignment = Alignment(horizontal='center', vertical='center')
    ws['A4'].fill  = fill('D6EAF8')
    _med = Side(style='medium')
    for _c in range(1, 17):
        ws.cell(row=4, column=_c).border = Border(
            left=_med  if _c == 1  else Side(style=None),
            right=_med if _c == 16 else Side(style=None),
            top=_med, bottom=_med,
        )
    ws.row_dimensions[4].height = 28

    # ── En-têtes colonnes 2 niveaux (lignes 5-6) ─────────────────────────────
    ws.row_dimensions[5].height = 24
    ws.row_dimensions[6].height = 22

    simple_h = [
        (1,  'Code PAI'),
        (2,  'Programmes / Projets / Activités'),
        (3,  'Localisation'),
        (4,  'Poids (%)'),
        (5,  'Indicateurs'),
        (6,  "Période d'exécution"),
        (7,  'Struct. Resp.'),
        (8,  'Coût total (F CFA)'),
        (9,  'Taux exéc. physique (%)'),
        (16, 'Observations'),
    ]
    for col, h in simple_h:
        ws.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col)
        c = ws.cell(row=5, column=col, value=h)
        c.font = Font(bold=True, size=9); c.fill = fill('BDD7EE')
        c.alignment = center; c.border = bord
        ws.cell(row=6, column=col).border = bord

    # Groupe "Exécution financière" cols 10-15 ligne 5
    ws.merge_cells(start_row=5, start_column=10, end_row=5, end_column=15)
    c = ws.cell(row=5, column=10, value='EXÉCUTION FINANCIÈRE (F CFA)')
    c.font = Font(bold=True, size=9); c.fill = fill('C9DAF8')
    c.alignment = center; c.border = bord

    fin_sub = [
        (10, 'Montant engagé'),
        (11, 'Taux eng. (%)'),
        (12, 'Montant mandaté'),
        (13, 'Taux mand. (%)'),
        (14, 'Montant payé'),
        (15, 'Taux paiem. (%)'),
    ]
    for col, h in fin_sub:
        c = ws.cell(row=6, column=col, value=h)
        c.font = Font(bold=True, size=9); c.fill = fill('C9DAF8')
        c.alignment = center; c.border = bord

    r = 7

    def _periode(act):
        if act.periode_debut and act.periode_fin and act.periode_debut != act.periode_fin:
            return f'{act.periode_debut} – {act.periode_fin}'
        return act.periode_debut or act.periode_fin or ''

    def _pct_cell(ws, r, col, val):
        c = ws.cell(row=r, column=col, value=round(val, 1))
        c.number_format = '0.0"%"'
        return c

    for pd in data:
        pg = pd['programme']
        pn = pd['prog_num']
        ws.cell(row=r, column=1, value=str(pn))
        ws.cell(row=r, column=2, value=f'Programme {pn} : {pg.nom}')
        ws.cell(row=r, column=4, value=pd['poids'])
        ws.cell(row=r, column=8, value=pd['budget'])
        _pct_cell(ws, r, 9, pd['taux_phys'])
        ws.cell(row=r, column=10, value=pd['engage'])
        _pct_cell(ws, r, 11, pd['taux_eng'])
        ws.cell(row=r, column=12, value=pd['mandate'])
        _pct_cell(ws, r, 13, pd['taux_mand'])
        ws.cell(row=r, column=14, value=pd['paye'])
        _pct_cell(ws, r, 15, pd['taux_pay'])
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
        ws.merge_cells(start_row=r, start_column=16, end_row=r, end_column=16)
        for col in range(1, 17):
            c = ws.cell(row=r, column=col)
            c.fill = fill('C8DCEA'); c.font = Font(bold=True, size=9); c.border = bord
        ws.cell(row=r, column=1).alignment = center
        ws.cell(row=r, column=2).alignment = left
        for col in [4, 8, 9, 10, 11, 12, 13, 14, 15]:
            ws.cell(row=r, column=col).alignment = center
        for col in [8, 10, 12, 14]:
            ws.cell(row=r, column=col).number_format = '#,##0'
        r += 1

        for pjd in pd['projets']:
            pj  = pjd['projet']
            pjn = pjd['proj_num']
            ws.cell(row=r, column=1, value=f'{pn}.{pjn}')
            ws.cell(row=r, column=2, value=f'Projet {pn}.{pjn} : {pj.nom}')
            ws.cell(row=r, column=4, value=pjd['poids'])
            ws.cell(row=r, column=8, value=pjd['budget'])
            _pct_cell(ws, r, 9, pjd['taux_phys'])
            ws.cell(row=r, column=10, value=pjd['engage'])
            _pct_cell(ws, r, 11, pjd['taux_eng'])
            ws.cell(row=r, column=12, value=pjd['mandate'])
            _pct_cell(ws, r, 13, pjd['taux_mand'])
            ws.cell(row=r, column=14, value=pjd['paye'])
            _pct_cell(ws, r, 15, pjd['taux_pay'])
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
            for col in range(1, 17):
                c = ws.cell(row=r, column=col)
                c.fill = fill('D4E6D4'); c.font = Font(bold=True, size=9); c.border = bord
            ws.cell(row=r, column=1).alignment = center
            ws.cell(row=r, column=2).alignment = left
            for col in [4, 8, 9, 10, 11, 12, 13, 14, 15]:
                ws.cell(row=r, column=col).alignment = center
            for col in [8, 10, 12, 14]:
                ws.cell(row=r, column=col).number_format = '#,##0'
            r += 1

            for act_i, ad in enumerate(pjd['activites'], 1):
                act   = ad['activite']
                extra = act.pai_extra
                row_data = [
                    f'{pn}.{pjn}.{act_i}',
                    act.nom,
                    extra.localisation if extra and extra.localisation else '',
                    ad['poids'],
                    extra.indicateurs if extra and extra.indicateurs else '',
                    _periode(act),
                    act.direction_responsable.code if act.direction_responsable else '',
                    ad['budget'],
                    round(ad['taux_phys'], 1),
                    ad['engage'],
                    round(ad['taux_eng'],  1),
                    ad['mandate'],
                    round(ad['taux_mand'], 1),
                    ad['paye'],
                    round(ad['taux_pay'],  1),
                    ad['obs'],
                ]
                for col, val in enumerate(row_data, 1):
                    c = ws.cell(row=r, column=col, value=val)
                    c.fill = fill('FFFFFF'); c.font = Font(size=9); c.border = bord
                    c.alignment = left if col == 2 else center
                    if col in (8, 10, 12, 14):
                        c.number_format = '#,##0'
                    if col in (9, 11, 13, 15):
                        c.number_format = '0.0"%"'
                r += 1

    # Ligne totaux
    ws.cell(row=r, column=1, value='TOTAL')
    ws.cell(row=r, column=2, value='TOTAL GÉNÉRAL')
    ws.cell(row=r, column=4, value=100)
    ws.cell(row=r, column=8, value=totaux['budget'])
    _pct_cell(ws, r, 9, totaux['taux_phys'])
    ws.cell(row=r, column=10, value=totaux['engage'])
    _pct_cell(ws, r, 11, totaux['taux_eng'])
    ws.cell(row=r, column=12, value=totaux['mandate'])
    _pct_cell(ws, r, 13, totaux['taux_mand'])
    ws.cell(row=r, column=14, value=totaux['paye'])
    _pct_cell(ws, r, 15, totaux['taux_pay'])
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
    for col in range(1, 17):
        c = ws.cell(row=r, column=col)
        c.fill = fill('1E3A5F'); c.font = Font(bold=True, size=9, color='FFFFFF'); c.border = bord
        c.alignment = center
    for col in [8, 10, 12, 14]:
        ws.cell(row=r, column=col).number_format = '#,##0'

    # Ligne pied : date impression + DDLP
    r += 1
    from datetime import datetime as _dt
    _now = _dt.now()
    _jours = ['dimanche','lundi','mardi','mercredi','jeudi','vendredi','samedi']
    _mois  = ['janvier','février','mars','avril','mai','juin','juillet','août','septembre','octobre','novembre','décembre']
    _date_str = f"{_jours[_now.weekday()]} {_now.day} {_mois[_now.month-1]} {_now.year} à {_now.strftime('%Hh%M')}"
    _dc = ws.cell(row=r, column=1, value=f'Édité le {_date_str}')
    _dc.font = Font(italic=True, size=8, color='444444')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    _dc.alignment = Alignment(horizontal='left', vertical='center')
    _ddr = ws.cell(row=r, column=9, value='Direction du Développement Local et de la Planification (DDLP)')
    _ddr.font = Font(bold=True, size=8)
    ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=16)
    _ddr.alignment = Alignment(horizontal='right', vertical='center')

    # Largeurs colonnes
    col_widths = [8, 34, 14, 7, 22, 12, 10, 16, 10, 16, 10, 16, 10, 16, 10, 28]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f'PEI_{annee.annee}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@pei_bp.route('/edit/<int:activite_id>', methods=['POST'])
@login_required
def edit_montants(activite_id):
    if not _peut_editer():
        abort(403)

    annee = get_annee()
    if not annee:
        abort(404)

    pei = PeiActivite.query.filter_by(activite_id=activite_id).first()
    if not pei:
        abort(404)

    budget = pei.activite.budget_total

    try:
        engage  = float(request.form.get('montant_engage',  0) or 0)
        mandate = float(request.form.get('montant_mandate', 0) or 0)
        paye    = float(request.form.get('montant_paye',    0) or 0)
    except ValueError:
        flash("Valeurs invalides.", 'danger')
        return redirect(url_for('pei.index'))

    # Contraintes
    if engage > budget:
        flash(f"Le montant engagé ({engage:,.0f}) ne peut pas dépasser le coût total ({budget:,.0f}).", 'warning')
        return redirect(url_for('pei.index', go=f'act-{activite_id}'))
    if mandate > engage:
        flash(f"Le montant mandaté ({mandate:,.0f}) ne peut pas dépasser le montant engagé ({engage:,.0f}).", 'warning')
        return redirect(url_for('pei.index', go=f'act-{activite_id}'))
    if paye > engage:
        flash(f"Le montant payé ({paye:,.0f}) ne peut pas dépasser le montant engagé ({engage:,.0f}).", 'warning')
        return redirect(url_for('pei.index', go=f'act-{activite_id}'))

    pei.montant_engage  = engage
    pei.montant_mandate = mandate
    pei.montant_paye    = paye
    db.session.commit()

    log_audit('pei_edit',
              f"Activité {pei.activite.code} — Engagé {engage:,.0f} / "
              f"Mandaté {mandate:,.0f} / Payé {paye:,.0f} F CFA")

    return redirect(url_for('pei.index', go=f'act-{activite_id}'))
